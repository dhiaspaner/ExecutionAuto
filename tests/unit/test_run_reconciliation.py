"""End-to-end tests for the interactive runner.

The wizard is driven from scripted keyboard input and both databases are fake,
so the whole path — fifteen questions, two connections, workbook read, guard,
execute, compare, write — runs with no database and no credential.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
import run_reconciliation
from openpyxl import Workbook, load_workbook

from migration_reconciliation.database.live import LiveExecutorFactory
from migration_reconciliation.models import DatabaseType
from migration_reconciliation.workbook.inline import build_inline_schema
from tests.drivers import FakeConnection, FakeOracleDb, FakePyodbc, odbc_error

SHEET = "Payments"
SOURCE_SECRET = "src-secret-not-real"
TARGET_SECRET = "tgt-secret-not-real"

#: Exactly the layout the runner documents: three inputs, five results.
HEADERS = [
    "ID",
    "Source SQL",
    "Target SQL",
    "Source Results",
    "Target Results",
    "Variance",
    "Status",
    "Remarks",
]

CASES = [
    ("TC-001", "SELECT COUNT(*) FROM PAYMENTS", "SELECT COUNT(*) FROM dbo.Payments"),
    ("TC-002", "SELECT SUM(amt) FROM PAYMENTS", "SELECT SUM(amt) FROM dbo.Payments"),
    ("TC-003", "SELECT COUNT(*) FROM FINES", "SELECT COUNT(*) FROM dbo.Fines"),
]


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    """A workbook in the runner's fixed column format, plus an unrelated sheet."""
    book = Workbook()
    notes = book.active
    notes.title = "Notes"
    notes["A1"] = "This sheet must survive the run untouched."

    sheet = book.create_sheet(SHEET)
    sheet.append(HEADERS)
    for case in CASES:
        sheet.append([*case, None, None, None, None, None])

    path = tmp_path / "payments_domain.xlsx"
    book.save(path)
    book.close()
    return path


def answers_for(
    workbook: Path,
    *,
    sheet: str | None = SHEET,
    source_type: str = "sqlserver",
    source_auth: str = "password",
    target_auth: str = "password",
) -> list[str]:
    """The keyboard answers for one run. ``sheet=None`` when it is not asked for."""
    common = [str(workbook)]
    if sheet is not None:
        common.append(sheet)
    common += [source_type, "legacy-host", "", "LegacyDb"]
    if source_type == "sqlserver":
        common.append("y")  # trust the source certificate
        common.append(source_auth)
    if source_type == "oracle" or source_auth == "password":
        common.append("legacy_reader")
    common += ["new-host", "", "MigratedDb", "y", target_auth]
    if target_auth == "password":
        common.append("migrated_reader")
    return common


@pytest.fixture
def console(monkeypatch: pytest.MonkeyPatch) -> Any:
    """Drive the real wizard from a scripted list of keyboard answers."""

    class Console:
        def __init__(self) -> None:
            self.answers: list[str] = []
            self.secrets: list[str] = [SOURCE_SECRET, TARGET_SECRET]

        def script(self, answers: list[str]) -> None:
            self.answers = list(answers)

        def _input(self, _prompt: str = "") -> str:
            if not self.answers:
                raise EOFError
            return self.answers.pop(0)

        def _getpass(self, _prompt: str = "") -> str:
            if not self.secrets:
                raise EOFError
            return self.secrets.pop(0)

    console = Console()
    monkeypatch.setattr("builtins.input", console._input)
    monkeypatch.setattr("migration_reconciliation.wizard.getpass_module.getpass", console._getpass)
    return console


def use_drivers(
    monkeypatch: pytest.MonkeyPatch, source_driver: Any, target_driver: Any
) -> dict[str, Any]:
    """Make the runner build its factory against the fake drivers."""
    built: dict[str, Any] = {}

    def factory(source: Any, target: Any) -> LiveExecutorFactory:
        built["factory"] = LiveExecutorFactory(
            source, target, source_driver=source_driver, target_driver=target_driver
        )
        return built["factory"]

    monkeypatch.setattr(run_reconciliation, "LiveExecutorFactory", factory)
    return built


def scripted_connection(values: dict[str, Any]) -> FakeConnection:
    return FakeConnection(values_by_sql=values)


# -- the full path -------------------------------------------------------


def test_a_complete_run_writes_a_timestamped_result_workbook(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any, tmp_path: Path
) -> None:
    console.script(answers_for(workbook))
    source = scripted_connection({CASES[0][1]: 1500, CASES[1][1]: 98765, CASES[2][1]: 12})
    target = scripted_connection({CASES[0][2]: 1500, CASES[1][2]: 98765, CASES[2][2]: 11})
    use_drivers(monkeypatch, FakePyodbc(connection=source), FakePyodbc(connection=target))

    code = run_reconciliation.main([])

    out = capsys.readouterr().out
    assert code == run_reconciliation.EXIT_FAILURES  # TC-003 legitimately mismatches
    assert "2 passed, 1 failed, 0 errors, 0 skipped" in out

    results = list(tmp_path.glob("payments_domain_results_*.xlsx"))
    assert len(results) == 1

    written = load_workbook(results[0])
    sheet = written[SHEET]
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert sheet["D2"].value == 1500
    assert sheet["E2"].value == 1500
    assert sheet["G2"].value == "PASS"
    assert sheet["G4"].value == "FAIL"
    assert "Notes" in written.sheetnames
    assert written["Notes"]["A1"].value == "This sheet must survive the run untouched."
    written.close()


def test_the_input_workbook_is_never_modified(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    before = workbook.read_bytes()
    console.script(answers_for(workbook))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    run_reconciliation.main([])

    assert workbook.read_bytes() == before


def test_an_oracle_source_uses_the_oracle_driver(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    console.script(answers_for(workbook, source_type="oracle"))
    source_driver = FakeOracleDb(connection=scripted_connection({CASES[0][1]: 7}))
    target_driver = FakePyodbc(connection=scripted_connection({CASES[0][2]: 7}))
    use_drivers(monkeypatch, source_driver, target_driver)

    run_reconciliation.main(["--case", "TC-001"])

    assert source_driver.connect_kwargs["dsn"] == "legacy-host:1521/LegacyDb"
    assert "oracle" in capsys.readouterr().out


def test_both_connections_are_closed_when_the_run_finishes(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = FakeConnection()
    target = FakeConnection()
    console.script(answers_for(workbook))
    use_drivers(monkeypatch, FakePyodbc(connection=source), FakePyodbc(connection=target))

    run_reconciliation.main([])

    assert source.closed is True
    assert target.closed is True


def test_nothing_connects_before_every_question_is_answered(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Input runs out at the last password, so no connection may have been made."""
    console.script(answers_for(workbook))
    console.secrets = [SOURCE_SECRET]  # the target password never arrives
    source_driver = FakePyodbc()
    target_driver = FakePyodbc()
    use_drivers(monkeypatch, source_driver, target_driver)

    code = run_reconciliation.main([])

    assert code == run_reconciliation.EXIT_USAGE
    assert source_driver.connections == []
    assert target_driver.connections == []


# -- flags ---------------------------------------------------------------


def test_case_runs_only_the_named_test(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    console.script(answers_for(workbook))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    code = run_reconciliation.main(["--case", "TC-002"])

    out = capsys.readouterr().out
    assert code == run_reconciliation.EXIT_OK
    assert "1 passed, 0 failed, 0 errors, 2 skipped" in out


def test_limit_runs_a_pilot(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    console.script(answers_for(workbook))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    run_reconciliation.main(["--limit", "2"])

    assert "2 passed, 0 failed, 0 errors, 1 skipped" in capsys.readouterr().out


def test_output_dir_places_the_result_workbook(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    destination = tmp_path / "results"
    destination.mkdir()
    console.script(answers_for(workbook))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    run_reconciliation.main(["--output-dir", str(destination)])

    assert len(list(destination.glob("payments_domain_results_*.xlsx"))) == 1


def test_no_password_flag_exists() -> None:
    """The absence of this flag is a security guarantee, so it is tested."""
    parser = run_reconciliation.build_parser()
    flags = {action.option_strings[0] for action in parser._actions if action.option_strings}

    assert flags == {
        "-h",
        "--profile",
        "--schema",
        "--mode",
        "--on-syntax-error",
        "--case",
        "--limit",
        "--output-dir",
    }
    # The guarantee this test exists for: nothing here accepts a credential.
    assert not any(
        word in flag for flag in flags for word in ("password", "pwd", "secret", "token", "user")
    )
    assert not any("password" in flag for flag in flags)
    assert not any("server" in flag or "user" in flag for flag in flags)


# -- failures ------------------------------------------------------------


def test_an_unreachable_source_stops_the_run_with_a_sanitized_message(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    console.script(answers_for(workbook))
    failing = FakePyodbc(connect_error=odbc_error("08001", f"cannot reach; PWD={SOURCE_SECRET}"))
    use_drivers(monkeypatch, failing, FakePyodbc())

    code = run_reconciliation.main([])

    captured = capsys.readouterr()
    assert code == run_reconciliation.EXIT_USAGE
    assert "server could not be reached" in captured.err
    assert SOURCE_SECRET not in captured.err + captured.out


def test_a_query_the_server_will_not_compile_is_an_error_and_the_rest_still_run(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    """An execute run finds an invalid object name by running the query.

    There is no compile pass in execute mode, so the database reports the bad
    table when that test runs. It is recorded as ERROR and the two sound
    queries still run: one broken row costs one result, not all of them.
    """
    console.script(answers_for(workbook))
    source = scripted_connection(
        {
            CASES[0][1]: 1500,
            CASES[1][1]: odbc_error("42S02", "Invalid object name"),
            CASES[2][1]: 12,
        }
    )
    target = scripted_connection({CASES[0][2]: 1500, CASES[1][2]: 0, CASES[2][2]: 12})
    use_drivers(monkeypatch, FakePyodbc(connection=source), FakePyodbc(connection=target))

    code = run_reconciliation.main([])

    out = capsys.readouterr().out
    assert code == run_reconciliation.EXIT_FAILURES
    # No compile pass: the run starts executing straight away.
    assert "Validation phase" not in out
    assert "Execution phase" in out
    assert "Invalid object name" in out
    # The sound queries still ran.
    assert "2 passed" in out


def test_unsafe_sql_is_rejected_before_it_reaches_the_database(
    tmp_path: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    sheet.append(["TC-BAD", "DELETE FROM PAYMENTS", "SELECT 1", None, None, None, None, None])
    path = tmp_path / "unsafe.xlsx"
    book.save(path)
    book.close()

    console.script(answers_for(path))
    source_driver = FakePyodbc()
    use_drivers(monkeypatch, source_driver, FakePyodbc())

    code = run_reconciliation.main([])

    out = capsys.readouterr().out
    assert code == run_reconciliation.EXIT_FAILURES
    assert "ERROR" in out
    assert source_driver.connection is not None
    assert source_driver.connection.executed_sql == []


def test_a_sheet_without_result_columns_is_refused(
    tmp_path: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(["ID", "Source SQL", "Target SQL"])
    sheet.append(["TC-001", "SELECT 1", "SELECT 1"])
    path = tmp_path / "no_results.xlsx"
    book.save(path)
    book.close()

    console.script(answers_for(path))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    code = run_reconciliation.main([])

    assert code == run_reconciliation.EXIT_USAGE
    assert "nowhere to write the outcome" in capsys.readouterr().err


def test_a_missing_required_column_is_reported_before_any_query(
    tmp_path: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET
    sheet.append(["ID", "Target SQL", "Status", "Remarks"])
    sheet.append(["TC-001", "SELECT 1", None, None])
    path = tmp_path / "missing_column.xlsx"
    book.save(path)
    book.close()

    console.script(answers_for(path))
    source_driver = FakePyodbc()
    use_drivers(monkeypatch, source_driver, FakePyodbc())

    code = run_reconciliation.main([])

    assert code == run_reconciliation.EXIT_USAGE
    assert "Source SQL" in capsys.readouterr().err
    assert source_driver.connection is not None
    assert source_driver.connection.executed_sql == []


# -- the fixed layout ----------------------------------------------------


def test_the_wizard_answer_decides_the_source_engine_not_a_stray_column() -> None:
    schema = build_inline_schema(SHEET, DatabaseType.ORACLE)

    assert schema.field("source_type").default == "oracle"
    assert schema.field("source_type").header.startswith("__wizard__")


def test_optional_columns_are_honoured_when_present_and_defaulted_when_absent() -> None:
    schema = build_inline_schema(SHEET, DatabaseType.SQLSERVER)

    assert schema.field("comparison_rule").header == "Comparison Rule"
    assert schema.field("comparison_rule").default == "equal"
    assert schema.field("comparison_rule").required is False
    # 0 means no limit: the runner imposes no query timeout of its own.
    assert schema.field("timeout_seconds").default == 0


# -- Windows authentication and profiles, end to end ---------------------


def test_windows_authentication_needs_no_password_at_all(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    console.script(answers_for(workbook, source_auth="windows", target_auth="windows"))
    console.secrets = []  # nothing may ask for one
    source = FakeConnection()
    target = FakeConnection()
    use_drivers(monkeypatch, FakePyodbc(connection=source), FakePyodbc(connection=target))

    code = run_reconciliation.main([])

    assert code == run_reconciliation.EXIT_OK
    assert "Trusted_Connection=yes" in source.connection_string
    assert "UID=" not in source.connection_string
    assert "PWD=" not in source.connection_string
    assert "Trusted_Connection=yes" in target.connection_string


def test_password_authentication_still_sends_credentials(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch
) -> None:
    console.script(answers_for(workbook))
    source = FakeConnection()
    use_drivers(monkeypatch, FakePyodbc(connection=source), FakePyodbc())

    run_reconciliation.main([])

    assert "UID=legacy_reader" in source.connection_string
    assert f"PWD={SOURCE_SECRET}" in source.connection_string
    assert "Trusted_Connection" not in source.connection_string


def test_the_connection_summary_names_windows_authentication(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any
) -> None:
    console.script(answers_for(workbook, source_auth="windows", target_auth="windows"))
    console.secrets = []
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    run_reconciliation.main([])

    assert "(Windows authentication)" in capsys.readouterr().out


def write_profile(path: Path, body: str) -> Path:
    path.write_text(body, encoding="utf-8")
    return path


def test_a_profile_answers_the_questions_through_the_real_entry_point(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: Any
) -> None:
    profile = write_profile(
        tmp_path / "profile.toml",
        f'''
version = "1.0"

[workbook]
path = "{workbook.as_posix()}"
sheet = "{SHEET}"

[source]
type = "sqlserver"
server = "legacy-host"
port = 1433
database = "LegacyDb"
trust_server_certificate = true
authentication = "windows"

[target]
server = "new-host"
port = 1433
database = "MigratedDb"
trust_server_certificate = false
authentication = "windows"
''',
    )
    console.script([])  # every question is answered by the file
    console.secrets = []
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    code = run_reconciliation.main(["--profile", str(profile)])

    out = capsys.readouterr().out
    assert code == run_reconciliation.EXIT_OK
    assert "[profile] Sheet: Payments" in out
    assert "3 passed" in out


def test_a_profile_still_prompts_for_the_password(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    profile = write_profile(
        tmp_path / "profile.toml",
        f'''
version = "1.0"

[workbook]
path = "{workbook.as_posix()}"
sheet = "{SHEET}"

[source]
type = "sqlserver"
server = "legacy-host"
port = 1433
database = "LegacyDb"
trust_server_certificate = true
authentication = "password"
username = "legacy_reader"

[target]
server = "new-host"
port = 1433
database = "MigratedDb"
trust_server_certificate = true
authentication = "password"
username = "svc_recon"
''',
    )
    console.script([])
    source = FakeConnection()
    use_drivers(monkeypatch, FakePyodbc(connection=source), FakePyodbc())

    code = run_reconciliation.main(["--profile", str(profile)])

    assert code == run_reconciliation.EXIT_OK
    assert f"PWD={SOURCE_SECRET}" in source.connection_string
    assert console.secrets == []  # both hidden prompts were consumed


def test_a_profile_containing_a_password_is_refused(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: Any
) -> None:
    profile = write_profile(
        tmp_path / "leaky.toml",
        f'''
version = "1.0"

[workbook]
path = "{workbook.as_posix()}"

[source]
type = "sqlserver"
server = "legacy-host"
database = "LegacyDb"
username = "legacy_reader"
password = "hunter2"
''',
    )
    console.script([])
    driver = FakePyodbc()
    use_drivers(monkeypatch, driver, FakePyodbc())

    code = run_reconciliation.main(["--profile", str(profile)])

    captured = capsys.readouterr()
    assert code == run_reconciliation.EXIT_USAGE
    assert "never contain a password" in captured.err
    assert "hunter2" not in captured.err + captured.out
    assert driver.connections == []


def test_a_missing_profile_file_is_reported(
    console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: Any
) -> None:
    console.script([])
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    code = run_reconciliation.main(["--profile", str(tmp_path / "nope.toml")])

    assert code == run_reconciliation.EXIT_USAGE
    assert "Cannot read profile" in capsys.readouterr().err


# -- the workbook schema DSL ---------------------------------------------

DSL_SHEET = "Definitions"
DSL_SECOND_SHEET = "Definitions Q4"

#: Deliberately unlike the runner's fixed layout: a title band above the
#: headers, and not one header text the runner knows by heart.
DSL_HEADERS = [
    "Case",
    "Run?",
    "Source Statement",
    "Target Statement",
    "Src Value",
    "Tgt Value",
    "Delta",
    "Outcome",
    "Notes",
]

DSL_HEADER_ROW = 3
DSL_FIRST_DATA_ROW = 4

SCHEMA_TOML = """\
schema_version = "1.0"
profile_name = "dsl-under-test"

[workbook]
test_case_sheet = "{sheet}"
header_row = {header_row}
first_data_row = {first_data_row}
preserve_other_sheets = true
output_filename_pattern = "{{input_stem}}_results_{{timestamp}}.xlsx"

[fields.test_case_id]
header = "Case"
type = "string"
required = true
read = true

[fields.enabled]
header = "Run?"
type = "boolean"
required = false
default = true
read = true

[fields.source_sql]
header = "Source Statement"
type = "sql"
required = true
read = true

[fields.target_sql]
header = "Target Statement"
type = "sql"
required = true
read = true

[fields.execution_scope]
header = "Not A Column: Execution Scope"
type = "enum"
allowed_values = ["SOURCE_TARGET", "SOURCE_ONLY", "TARGET_ONLY"]
required = false
default = "SOURCE_TARGET"
read = true

[fields.source_type]
header = "Not A Column: Source Type"
type = "enum"
allowed_values = ["sqlserver", "oracle"]
required = false
default = "sqlserver"
read = true

[fields.source_connection]
header = "Not A Column: Source Connection"
type = "string"
required = false
default = "SOURCE"
read = true

[fields.target_connection]
header = "Not A Column: Target Connection"
type = "string"
required = false
default = "TARGET"
read = true

[fields.source_result]
header = "Src Value"
type = "scalar"
required = false
write = true

[fields.target_result]
header = "Tgt Value"
type = "scalar"
required = false
write = true

[fields.variance]
header = "Delta"
type = "decimal"
required = false
write = true

[fields.status]
header = "Outcome"
type = "enum"
allowed_values = ["PASS", "FAIL", "ERROR", "SKIPPED"]
required = false
write = true

[fields.remarks]
header = "Notes"
type = "string"
required = false
write = true

[fields.executed_at]
header = "Not A Column: Executed At"
type = "datetime"
required = false
write = true

[fields.run_id]
header = "Not A Column: Run ID"
type = "string"
required = false
write = true

[fields.duration_ms]
header = "Not A Column: Duration"
type = "integer"
required = false
write = true

[fields.error_side]
header = "Not A Column: Error Side"
type = "enum"
allowed_values = ["SOURCE", "TARGET", "COMPARISON", "WORKBOOK", ""]
required = false
write = true
"""


def dsl_sheet(book: Workbook, title: str, cases: list[tuple[str, str, str]]) -> None:
    """One sheet in the DSL layout: banner, blank line, headers, then rows."""
    sheet = book.create_sheet(title)
    sheet.append([f"{title} - definitions above, results to the right"] * len(DSL_HEADERS))
    sheet.append([])
    sheet.append(DSL_HEADERS)
    for case in cases:
        sheet.append([case[0], "Yes", case[1], case[2]])


@pytest.fixture
def dsl_workbook(tmp_path: Path) -> Path:
    """A workbook no fixed-layout reader can read: the headers sit on row 3."""
    book = Workbook()
    book.remove(book.active)
    dsl_sheet(book, DSL_SHEET, CASES)
    dsl_sheet(book, DSL_SECOND_SHEET, CASES[:1])

    path = tmp_path / "payments_dsl.xlsx"
    book.save(path)
    book.close()
    return path


def schema_file(
    tmp_path: Path,
    *,
    sheet: str = DSL_SHEET,
    header_row: int = DSL_HEADER_ROW,
    first_data_row: int = DSL_FIRST_DATA_ROW,
) -> Path:
    path = tmp_path / "schema.toml"
    path.write_text(
        SCHEMA_TOML.format(sheet=sheet, header_row=header_row, first_data_row=first_data_row),
        encoding="utf-8",
    )
    return path


def profile_file(tmp_path: Path, workbook_table: str) -> Path:
    path = tmp_path / "profile.toml"
    path.write_text(f'version = "1.0"\n\n[workbook]\n{workbook_table}\n', encoding="utf-8")
    return path


def test_a_schema_decides_the_sheet_the_header_row_and_the_headers(
    dsl_workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any, tmp_path: Path
) -> None:
    console.script(answers_for(dsl_workbook, sheet=None))
    source = scripted_connection({CASES[0][1]: 1500, CASES[1][1]: 98765, CASES[2][1]: 12})
    target = scripted_connection({CASES[0][2]: 1500, CASES[1][2]: 98765, CASES[2][2]: 11})
    use_drivers(monkeypatch, FakePyodbc(connection=source), FakePyodbc(connection=target))

    code = run_reconciliation.main(["--schema", str(schema_file(tmp_path))])

    out = capsys.readouterr().out
    assert code == run_reconciliation.EXIT_FAILURES  # TC-003 legitimately mismatches
    assert "2 passed, 1 failed, 0 errors, 0 skipped" in out
    assert "headers on row 3" in out

    written = load_workbook(next(tmp_path.glob("payments_dsl_results_*.xlsx")))
    sheet = written[DSL_SHEET]
    assert [cell.value for cell in sheet[DSL_HEADER_ROW]] == DSL_HEADERS
    assert sheet["E4"].value == 1500  # Src Value
    assert sheet["F4"].value == 1500  # Tgt Value
    assert sheet["H4"].value == "PASS"  # Outcome
    assert sheet["H6"].value == "FAIL"
    written.close()


def test_a_schema_names_the_columns_before_the_first_question(
    dsl_workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any, tmp_path: Path
) -> None:
    console.script(answers_for(dsl_workbook, sheet=None))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    run_reconciliation.main(["--schema", str(schema_file(tmp_path)), "--limit", "1"])

    banner = capsys.readouterr().out.split("Workbook")[0]
    assert "Case, Source Statement, Target Statement" in banner
    assert "Src Value, Tgt Value, Delta, Outcome, Notes" in banner
    assert "dsl-under-test" in banner
    # The fixed layout's own header text must not be claimed for this workbook.
    assert "Source SQL" not in banner


def test_the_profile_can_name_the_schema(
    dsl_workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    profile = profile_file(tmp_path, f'schema = "{schema_file(tmp_path).as_posix()}"')
    console.script(answers_for(dsl_workbook, sheet=None))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    code = run_reconciliation.main(["--profile", str(profile)])

    assert code == run_reconciliation.EXIT_OK
    assert list(tmp_path.glob("payments_dsl_results_*.xlsx"))


def test_the_flag_overrides_the_schema_the_profile_names(
    dsl_workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    unusable = tmp_path / "wrong.toml"
    unusable.write_text('schema_version = "9.9"\n', encoding="utf-8")
    profile = profile_file(tmp_path, f'schema = "{unusable.as_posix()}"')
    console.script(answers_for(dsl_workbook, sheet=None))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    code = run_reconciliation.main(
        ["--profile", str(profile), "--schema", str(schema_file(tmp_path))]
    )

    assert code == run_reconciliation.EXIT_OK


def test_the_answered_sheet_wins_over_the_one_the_schema_declares(
    dsl_workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    before = load_workbook(dsl_workbook)
    untouched = [cell.value for cell in before[DSL_SHEET][DSL_FIRST_DATA_ROW]]
    before.close()

    profile = profile_file(
        tmp_path,
        f'sheet = "{DSL_SECOND_SHEET}"\nschema = "{schema_file(tmp_path).as_posix()}"',
    )
    console.script(answers_for(dsl_workbook, sheet=None))
    connection = scripted_connection({CASES[0][1]: 7, CASES[0][2]: 7})
    use_drivers(monkeypatch, FakePyodbc(connection=connection), FakePyodbc(connection=connection))

    code = run_reconciliation.main(["--profile", str(profile)])

    assert code == run_reconciliation.EXIT_OK
    written = load_workbook(next(tmp_path.glob("payments_dsl_results_*.xlsx")))
    assert written[DSL_SECOND_SHEET]["H4"].value == "PASS"
    assert [cell.value for cell in written[DSL_SHEET][DSL_FIRST_DATA_ROW]] == untouched
    written.close()


def test_an_unreadable_schema_stops_the_run_before_a_single_question(
    dsl_workbook: Path, console: Any, capsys: Any, tmp_path: Path
) -> None:
    broken = tmp_path / "broken.toml"
    broken.write_text('schema_version = "9.9"\nprofile_name = "nope"\n', encoding="utf-8")
    console.script(answers_for(dsl_workbook, sheet=None))

    code = run_reconciliation.main(["--schema", str(broken)])

    assert code == run_reconciliation.EXIT_USAGE
    assert "unsupported schema_version" in capsys.readouterr().err
    # Nothing was asked, so every scripted answer is still waiting.
    assert len(console.answers) == len(answers_for(dsl_workbook, sheet=None))


def test_without_a_schema_the_runners_own_layout_still_applies(
    workbook: Path, console: Any, monkeypatch: pytest.MonkeyPatch, capsys: Any, tmp_path: Path
) -> None:
    console.script(answers_for(workbook))
    use_drivers(monkeypatch, FakePyodbc(), FakePyodbc())

    code = run_reconciliation.main([])

    assert code == run_reconciliation.EXIT_OK
    assert "ID, Source SQL, Target SQL" in capsys.readouterr().out
    assert list(tmp_path.glob("payments_domain_results_*.xlsx"))
