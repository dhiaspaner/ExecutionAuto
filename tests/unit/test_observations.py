"""Observation rules: deterministic selection, safe rendering, no verdicts."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from migration_reconciliation.errors import ConfigurationError
from migration_reconciliation.models import Platform, TestStatus
from migration_reconciliation.workbook.observations import (
    BUILT_IN_OBSERVATIONS,
    ObservationRule,
    ObservationRules,
    read_observation_rules,
    render_template,
)


def rule(
    status: TestStatus,
    platform: Platform | None,
    error_code: str | None,
    template: str,
    *,
    row: int = 2,
    enabled: bool = True,
) -> ObservationRule:
    return ObservationRule(
        row_number=row,
        status=status,
        platform=platform,
        error_code=error_code,
        template=template,
        enabled=enabled,
    )


def rules_of(path: Path) -> ObservationRules:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return read_observation_rules(workbook)
    finally:
        workbook.close()


# -- selection ---------------------------------------------------------------


def test_an_exact_match_wins_over_every_wildcard() -> None:
    table = ObservationRules(
        (
            rule(TestStatus.ERROR, None, None, "any error", row=2),
            rule(TestStatus.ERROR, Platform.TARGET, None, "target error", row=3),
            rule(TestStatus.ERROR, None, "QUERY_TIMEOUT", "a timeout somewhere", row=4),
            rule(TestStatus.ERROR, Platform.TARGET, "QUERY_TIMEOUT", "exact", row=5),
        )
    )

    chosen = table.select(TestStatus.ERROR, Platform.TARGET, "QUERY_TIMEOUT")

    assert chosen is not None
    assert chosen.template == "exact"


def test_selection_falls_back_in_the_documented_order() -> None:
    table = ObservationRules(
        (
            rule(TestStatus.ERROR, None, None, "status only", row=2),
            rule(TestStatus.ERROR, Platform.TARGET, None, "platform match", row=3),
            rule(TestStatus.ERROR, None, "QUERY_TIMEOUT", "code match", row=4),
        )
    )

    assert table.select(TestStatus.ERROR, Platform.TARGET, "OTHER").template == "platform match"
    assert table.select(TestStatus.ERROR, Platform.SOURCE, "QUERY_TIMEOUT").template == "code match"
    assert table.select(TestStatus.ERROR, Platform.SOURCE, "OTHER").template == "status only"


def test_the_built_in_wording_is_used_when_no_rule_matches() -> None:
    table = ObservationRules((rule(TestStatus.PASS, None, None, "passed"),))

    assert (
        table.template_for(TestStatus.FAIL, Platform.COMPARISON, "VALUE_MISMATCH")
        == (BUILT_IN_OBSERVATIONS[TestStatus.FAIL])
    )


def test_a_disabled_rule_never_matches() -> None:
    table = ObservationRules(
        (
            rule(TestStatus.PASS, None, None, "disabled wording", enabled=False),
            rule(TestStatus.PASS, None, None, "live wording", row=3),
        )
    )

    assert table.select(TestStatus.PASS, Platform.NONE, None).template == "live wording"


# -- the sheet ---------------------------------------------------------------


def test_the_shipped_rules_load(make_recon_workbook: Any, test_row: Any) -> None:
    table = rules_of(make_recon_workbook([test_row()]))

    assert table.rules
    chosen = table.select(TestStatus.ERROR, Platform.TARGET, "QUERY_TIMEOUT")
    assert chosen is not None
    assert "{timeout_seconds}" in chosen.template


def test_two_enabled_rules_with_one_key_are_rejected(
    make_recon_workbook: Any, test_row: Any
) -> None:
    """Without a priority column, a tie would make the wording depend on row order."""
    duplicate = [
        {
            "Enabled": "Yes",
            "Status": "FAIL",
            "Platform": "COMPARISON",
            "Error_Code": "VALUE_MISMATCH",
            "Observation_Template": "first",
        },
        {
            "Enabled": "Yes",
            "Status": "FAIL",
            "Platform": "COMPARISON",
            "Error_Code": "VALUE_MISMATCH",
            "Observation_Template": "second",
        },
    ]
    path = make_recon_workbook([test_row()], observation_rules=duplicate)

    with pytest.raises(ConfigurationError, match="deterministic"):
        rules_of(path)


def test_a_duplicate_key_is_allowed_when_one_side_is_disabled(
    make_recon_workbook: Any, test_row: Any
) -> None:
    pair = [
        {
            "Enabled": "No",
            "Status": "FAIL",
            "Platform": "COMPARISON",
            "Error_Code": "VALUE_MISMATCH",
            "Observation_Template": "retired",
        },
        {
            "Enabled": "Yes",
            "Status": "FAIL",
            "Platform": "COMPARISON",
            "Error_Code": "VALUE_MISMATCH",
            "Observation_Template": "current",
        },
    ]
    table = rules_of(make_recon_workbook([test_row()], observation_rules=pair))

    assert table.select(TestStatus.FAIL, Platform.COMPARISON, "VALUE_MISMATCH").template == (
        "current"
    )


def test_an_unknown_status_is_refused(make_recon_workbook: Any, test_row: Any) -> None:
    bad = [{"Enabled": "Yes", "Status": "NEARLY", "Observation_Template": "x"}]

    with pytest.raises(ConfigurationError, match="Status"):
        rules_of(make_recon_workbook([test_row()], observation_rules=bad))


def test_an_unknown_platform_is_refused_rather_than_treated_as_any(
    make_recon_workbook: Any, test_row: Any
) -> None:
    bad = [
        {
            "Enabled": "Yes",
            "Status": "FAIL",
            "Platform": "DATABSE",
            "Observation_Template": "x",
        }
    ]

    with pytest.raises(ConfigurationError, match="Platform"):
        rules_of(make_recon_workbook([test_row()], observation_rules=bad))


def test_an_unknown_placeholder_is_refused(make_recon_workbook: Any, test_row: Any) -> None:
    bad = [
        {
            "Enabled": "Yes",
            "Status": "PASS",
            "Observation_Template": "{test_id} took {wall_clock}",
        }
    ]

    with pytest.raises(ConfigurationError, match="wall_clock"):
        rules_of(make_recon_workbook([test_row()], observation_rules=bad))


# -- rendering ---------------------------------------------------------------


def test_only_known_placeholders_are_replaced() -> None:
    rendered = render_template(
        "{test_id}: {actual_value} of {expected_value}",
        {"test_id": "TC-1", "actual_value": "9", "expected_value": "10"},
        max_length=200,
    )

    assert rendered == "TC-1: 9 of 10"


def test_a_missing_value_renders_as_nothing_rather_than_an_error() -> None:
    assert render_template("{test_id}|{variance}", {"test_id": "TC-1"}, max_length=200) == "TC-1|"


def test_rendering_never_evaluates_anything() -> None:
    """No attribute access, no indexing, no expressions: substitution only."""
    rendered = render_template(
        "{test_id.__class__} {0} {}",
        {"test_id": "TC-1"},
        max_length=200,
    )

    assert rendered == "{test_id.__class__} {0} {}"


def test_observations_are_truncated_to_the_configured_length() -> None:
    rendered = render_template("{error_detail}", {"error_detail": "x" * 500}, max_length=60)

    assert len(rendered) <= 60


def test_observations_are_sanitized_before_they_reach_a_cell() -> None:
    rendered = render_template(
        "{error_detail}",
        {"error_detail": "login failed for PWD=hunter2;DATABASE=x"},
        max_length=300,
    )

    assert "hunter2" not in rendered
