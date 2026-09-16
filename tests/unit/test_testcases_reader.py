"""Reading and validating the ``Test Cases`` sheet."""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from migration_reconciliation.errors import WorkbookError
from migration_reconciliation.models import (
    ComparisonType,
    ErrorCode,
    ExecutionScope,
    ResultType,
    TestStatus,
)
from migration_reconciliation.workbook.testcases import (
    DefinitionProblem,
    TestCaseSheet,
    read_test_cases,
)


def read(path: Path) -> TestCaseSheet:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return read_test_cases(workbook, sheet_name="Test Cases")
    finally:
        workbook.close()


def only_problem(sheet: TestCaseSheet) -> DefinitionProblem:
    assert len(sheet.problems) == 1, sheet.problems
    return sheet.problems[0]


def source_only(test_row: Any, **overrides: Any) -> Mapping[str, Any]:
    row = test_row(
        "TC-SRC",
        Execution_Scope="SOURCE_ONLY",
        Comparison_Type="EXPECTED_ZERO",
        Target_Profile_Section=None,
        Target_SQL=None,
    )
    return {**row, **overrides}


def target_only(test_row: Any, **overrides: Any) -> Mapping[str, Any]:
    row = test_row(
        "TC-TGT",
        Execution_Scope="TARGET_ONLY",
        Comparison_Type="EXPECTED_ZERO",
        Source_Profile_Section=None,
        Source_SQL=None,
    )
    return {**row, **overrides}


# -- valid definitions -------------------------------------------------------


def test_a_source_target_test_needs_both_sides(make_recon_workbook: Any, test_row: Any) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001")]))

    assert not sheet.problems
    definition = sheet.enabled[0]
    assert definition.scope is ExecutionScope.SOURCE_TARGET
    assert definition.comparison is ComparisonType.EQUAL
    assert definition.result_type is ResultType.INTEGER
    assert definition.sections() == ("source", "target")


def test_a_source_only_test_names_only_the_source(make_recon_workbook: Any, test_row: Any) -> None:
    sheet = read(make_recon_workbook([source_only(test_row)]))

    assert not sheet.problems
    definition = sheet.enabled[0]
    assert definition.scope is ExecutionScope.SOURCE_ONLY
    assert definition.target_sql == ""
    assert definition.sections() == ("source",)


def test_a_target_only_test_names_only_the_target(make_recon_workbook: Any, test_row: Any) -> None:
    sheet = read(make_recon_workbook([target_only(test_row)]))

    assert not sheet.problems
    assert sheet.enabled[0].sections() == ("target",)


def test_headers_are_matched_by_name_not_position(make_recon_workbook: Any, test_row: Any) -> None:
    """The template may be re-ordered; nothing here depends on a column letter."""
    path = make_recon_workbook([test_row("TC-001")])
    sheet = read(path)

    assert sheet.column_of("Test_ID") != sheet.column_of("Status")
    assert sheet.column_of("Source_SQL") is not None


# -- disabled ----------------------------------------------------------------


def test_a_disabled_test_is_recorded_and_needs_no_connection(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sheet = read(make_recon_workbook([test_row("TC-OFF", Enabled="No")]))

    assert [d.test_id for d in sheet.disabled] == ["TC-OFF"]
    assert sheet.enabled == ()
    assert sheet.required_sections() == ()


def test_a_disabled_test_is_never_validated(make_recon_workbook: Any, test_row: Any) -> None:
    """A switched-off row must not be able to stop a run."""
    broken = test_row(
        "TC-OFF", Enabled="No", Execution_Scope="NONSENSE", Comparison_Type="INVENTED"
    )
    sheet = read(make_recon_workbook([broken]))

    assert not sheet.problems
    assert len(sheet.disabled) == 1


def test_enabled_must_be_yes_or_no(make_recon_workbook: Any, test_row: Any) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001", Enabled="maybe")]))

    assert only_problem(sheet).code is ErrorCode.INVALID_ENABLED


# -- identity ----------------------------------------------------------------


def test_a_missing_test_id_is_reported_against_its_row(
    make_recon_workbook: Any, test_row: Any
) -> None:
    row = test_row("TC-001")
    row.pop("Test_ID")
    sheet = read(make_recon_workbook([row]))

    assert only_problem(sheet).code is ErrorCode.MISSING_TEST_ID


def test_duplicate_ids_are_rejected(make_recon_workbook: Any, test_row: Any) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001"), test_row("TC-001")]))

    problem = only_problem(sheet)
    assert problem.code is ErrorCode.DUPLICATE_TEST_ID
    assert len(sheet.enabled) == 1


# -- controlled values -------------------------------------------------------


@pytest.mark.parametrize(
    ("column", "value", "code"),
    [
        ("Execution_Scope", "BOTH_SIDES", ErrorCode.INVALID_SCOPE),
        ("Comparison_Type", "ROUGHLY_EQUAL", ErrorCode.INVALID_COMPARISON_TYPE),
        ("Result_Type", "MONEY", ErrorCode.INVALID_RESULT_TYPE),
    ],
)
def test_controlled_columns_reject_invented_values(
    make_recon_workbook: Any, test_row: Any, column: str, value: str, code: ErrorCode
) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001", **{column: value})]))

    assert only_problem(sheet).code is code


def test_a_numeric_comparison_refuses_a_text_result_type(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sheet = read(
        make_recon_workbook(
            [test_row("TC-001", Comparison_Type="EXPECTED_ZERO", Result_Type="TEXT")]
        )
    )

    assert only_problem(sheet).code is ErrorCode.INVALID_RESULT_TYPE


def test_a_two_sided_comparison_refuses_a_one_sided_scope(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sheet = read(
        make_recon_workbook([source_only(test_row, Comparison_Type="EQUAL_ABS_TOLERANCE")])
    )

    problem = only_problem(sheet)
    assert problem.code is ErrorCode.INVALID_COMPARISON_TYPE
    assert "SOURCE_TARGET" in problem.message


@pytest.mark.parametrize("column", ["Absolute_Tolerance", "Percentage_Tolerance"])
@pytest.mark.parametrize("value", ["-1", "lots"])
def test_tolerances_must_be_numbers_of_zero_or_more(
    make_recon_workbook: Any, test_row: Any, column: str, value: str
) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001", **{column: value})]))

    assert only_problem(sheet).code is ErrorCode.INVALID_TOLERANCE


def test_tolerances_default_to_zero(make_recon_workbook: Any, test_row: Any) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001")]))

    assert sheet.enabled[0].absolute_tolerance == Decimal(0)
    assert sheet.enabled[0].percentage_tolerance == Decimal(0)


# -- scope rules -------------------------------------------------------------


def test_a_missing_required_query_is_a_config_error(
    make_recon_workbook: Any, test_row: Any
) -> None:
    row = test_row("TC-001")
    row.pop("Source_SQL")
    sheet = read(make_recon_workbook([row]))

    assert only_problem(sheet).code is ErrorCode.MISSING_SOURCE_SQL


def test_a_query_on_an_unused_side_is_refused(make_recon_workbook: Any, test_row: Any) -> None:
    """An ambiguous test is refused rather than half-run."""
    sheet = read(
        make_recon_workbook(
            [
                source_only(
                    test_row,
                    Target_SQL="SELECT COUNT(*) FROM dbo.Payments",
                    Target_Profile_Section="target",
                )
            ]
        )
    )

    problem = only_problem(sheet)
    assert problem.code in {
        ErrorCode.UNEXPECTED_TARGET_SQL,
        ErrorCode.UNEXPECTED_TARGET_SECTION,
    }


def test_a_missing_profile_section_is_a_config_error(
    make_recon_workbook: Any, test_row: Any
) -> None:
    row = test_row("TC-001")
    row.pop("Target_Profile_Section")
    sheet = read(make_recon_workbook([row]))

    assert only_problem(sheet).code is ErrorCode.MISSING_TARGET_SECTION


def test_a_section_the_profile_does_not_have_is_refused(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001", Source_Profile_Section="staging")]))

    problem = only_problem(sheet)
    assert problem.code is ErrorCode.MISSING_SOURCE_SECTION
    assert "staging" in problem.message


def test_section_names_tolerate_brackets_and_case(make_recon_workbook: Any, test_row: Any) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001", Source_Profile_Section="[SOURCE]")]))

    assert not sheet.problems
    assert sheet.enabled[0].source_section == "source"


# -- SQL safety --------------------------------------------------------------


@pytest.mark.parametrize(
    "sql",
    [
        "UPDATE dbo.Payments SET Status = 'X'",
        "SELECT 1; DROP TABLE dbo.Payments",
        "SELECT COUNT(*) FROM dbo.Payments -- fine\nDELETE FROM dbo.Payments",
        "EXEC sp_who",
    ],
)
def test_unsafe_sql_never_reaches_a_driver(
    make_recon_workbook: Any, test_row: Any, sql: str
) -> None:
    sheet = read(make_recon_workbook([test_row("TC-001", Target_SQL=sql)]))

    assert only_problem(sheet).code is ErrorCode.UNSAFE_SQL


def test_a_cte_that_ends_in_a_read_is_allowed(make_recon_workbook: Any, test_row: Any) -> None:
    cte = (
        "WITH recent AS (SELECT Id FROM dbo.Payments WHERE PaidOn > '2026-01-01') "
        "SELECT COUNT(*) FROM recent"
    )
    sheet = read(make_recon_workbook([test_row("TC-001", Target_SQL=cte)]))

    assert not sheet.problems


def test_cross_database_identifiers_are_left_exactly_as_written(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sql = "SELECT COUNT(*) FROM DXBPRODSQL02.dbo.Ledger"
    sheet = read(make_recon_workbook([test_row("TC-001", Source_SQL=sql)]))

    assert sheet.enabled[0].source_sql == sql


# -- expected values ---------------------------------------------------------


def test_a_comparison_that_needs_an_expectation_says_so(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sheet = read(make_recon_workbook([target_only(test_row, Comparison_Type="EXPECTED_EQUAL")]))

    assert only_problem(sheet).code is ErrorCode.MISSING_EXPECTED_VALUE


def test_an_expected_value_is_read_as_the_declared_result_type(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sheet = read(
        make_recon_workbook(
            [target_only(test_row, Comparison_Type="EXPECTED_EQUAL", Expected_Value="12")]
        )
    )

    assert sheet.enabled[0].expected_value == 12


def test_an_unconvertible_expected_value_is_caught_before_any_query(
    make_recon_workbook: Any, test_row: Any
) -> None:
    sheet = read(
        make_recon_workbook(
            [target_only(test_row, Comparison_Type="EXPECTED_EQUAL", Expected_Value="lots")]
        )
    )

    assert only_problem(sheet).code is ErrorCode.TYPE_CONVERSION_ERROR


# -- sheet structure ---------------------------------------------------------


def test_one_bad_row_never_stops_the_others(make_recon_workbook: Any, test_row: Any) -> None:
    rows: Sequence[Mapping[str, Any]] = [
        test_row("TC-001"),
        test_row("TC-002", Execution_Scope="SIDEWAYS"),
        test_row("TC-003"),
    ]
    sheet = read(make_recon_workbook(rows))

    assert [d.test_id for d in sheet.enabled] == ["TC-001", "TC-003"]
    assert only_problem(sheet).test_id == "TC-002"


def test_a_missing_required_column_stops_the_run(tmp_path: Path, test_row: Any) -> None:
    from migration_reconciliation.workbook.payments_template import build_workbook

    workbook = build_workbook([test_row("TC-001")])
    sheet = workbook["Test Cases"]
    for cell in sheet[1]:
        if cell.value == "Source_SQL":
            cell.value = "Src SQL"
    path = tmp_path / "broken.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(WorkbookError, match="Source_SQL"):
        read(path)


def test_a_sheet_with_nowhere_to_write_the_outcome_is_refused(
    tmp_path: Path, test_row: Any
) -> None:
    from migration_reconciliation.workbook.payments_template import build_workbook

    workbook = build_workbook([test_row("TC-001")])
    sheet = workbook["Test Cases"]
    for cell in sheet[1]:
        if cell.value == "Status":
            cell.value = "Outcome"
    path = tmp_path / "no-status.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(WorkbookError, match="Status"):
        read(path)


def test_legacy_columns_are_noticed_but_never_obeyed(tmp_path: Path, test_row: Any) -> None:
    from migration_reconciliation.workbook.payments_template import build_workbook

    workbook = build_workbook([test_row("TC-001")])
    sheet = workbook["Test Cases"]
    column = sheet.max_column + 1
    sheet.cell(row=1, column=column, value="Executor_Action")
    sheet.cell(row=2, column=column, value="DROP TABLE dbo.Payments")
    path = tmp_path / "legacy.xlsx"
    workbook.save(path)
    workbook.close()

    parsed = read(path)

    assert parsed.legacy_columns == ("Executor_Action",)
    assert not parsed.problems
    assert parsed.enabled[0].target_sql == "SELECT COUNT(*) FROM dbo.Payments"


def test_a_missing_sheet_is_reported_with_what_is_there(
    make_recon_workbook: Any, test_row: Any
) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        with pytest.raises(WorkbookError, match="Available sheets"):
            read_test_cases(workbook, sheet_name="Missing Sheet")
    finally:
        workbook.close()


def test_the_status_vocabulary_is_the_documented_one() -> None:
    assert {status.value for status in TestStatus} == {
        "PASS",
        "FAIL",
        "PROFILED",
        "VALIDATED",
        "ERROR",
        "SYNTAX ERROR",
        "BLOCKED",
        "CONFIG ERROR",
        "NOT EXECUTED",
        "DISABLED",
    }
