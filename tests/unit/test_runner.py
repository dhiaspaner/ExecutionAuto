"""End-to-end offline runs: execution, isolation, cleanup and selection."""

from __future__ import annotations

import os
from collections.abc import Callable
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from migration_reconciliation.database.base import QuerySide
from migration_reconciliation.database.factory import FakeExecutorFactory
from migration_reconciliation.errors import WorkbookError
from migration_reconciliation.models import ErrorSide, ExecutionStatus
from migration_reconciliation.runner import ReconciliationRunner, RunOptions


def _run(schema: Any, workbook: Path, book: Any, **option_overrides: Any):
    factory = FakeExecutorFactory(book)
    runner = ReconciliationRunner(schema, factory)
    summary = runner.run(workbook, RunOptions(**option_overrides))
    return summary, factory


def _by_id(summary: Any) -> dict[str, Any]:
    return {result.test_case_id: result for result in summary.results}


def test_passing_case_is_executed_and_recorded(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1500, "target_result": 1500})

    summary, _ = _run(schema, path, book)

    result = _by_id(summary)["TC-001"]
    assert result.status is ExecutionStatus.PASS
    assert result.source_result == 1500
    assert result.target_result == 1500
    assert result.variance == Decimal(0)
    assert result.error_side is ErrorSide.NONE
    assert result.run_id == summary.run_id
    assert summary.is_clean is True


def test_source_and_target_queries_are_sent_to_their_own_connections(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    _, factory = _run(schema, path, book)

    sent = {(e.connection_name, e.side): e.executed_sql for e in factory.created}
    assert sent[("LEGACY_ORACLE", QuerySide.SOURCE)] == ["SELECT COUNT(*) FROM PAYMENTS"]
    assert sent[("MIGRATED_SQLSERVER", QuerySide.TARGET)] == ["SELECT COUNT(*) FROM dbo.Payments"]


def test_failing_comparison_is_reported_without_stopping_the_run(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001"), case_row("TC-002")])
    book = make_results(
        {"test_case_id": "TC-001", "source_result": 10, "target_result": 7},
        {"test_case_id": "TC-002", "source_result": 5, "target_result": 5},
    )

    summary, _ = _run(schema, path, book)

    results = _by_id(summary)
    assert results["TC-001"].status is ExecutionStatus.FAIL
    assert results["TC-001"].variance == Decimal(3)
    assert results["TC-002"].status is ExecutionStatus.PASS
    assert summary.is_clean is False


def test_source_execution_error_is_isolated(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001"), case_row("TC-002")])
    book = make_results(
        {"test_case_id": "TC-001", "source_error": "ORA-00942: table or view does not exist"},
        {"test_case_id": "TC-002", "source_result": 5, "target_result": 5},
    )

    summary, _ = _run(schema, path, book)

    failed = _by_id(summary)["TC-001"]
    assert failed.status is ExecutionStatus.ERROR
    assert failed.error_side is ErrorSide.SOURCE
    assert failed.error_code == "SOURCE_EXEC_FAILED"
    assert "ORA-00942" in failed.remarks
    assert _by_id(summary)["TC-002"].status is ExecutionStatus.PASS


def test_target_execution_error_is_attributed_to_the_target(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results(
        {
            "test_case_id": "TC-001",
            "source_result": 10,
            "target_error": "Login timeout expired",
        }
    )

    summary, _ = _run(schema, path, book)

    result = _by_id(summary)["TC-001"]
    assert result.status is ExecutionStatus.ERROR
    assert result.error_side is ErrorSide.TARGET
    assert result.error_code == "TARGET_EXEC_FAILED"
    assert result.source_result == 10, "the source result is still recorded"


def test_error_messages_are_sanitized_before_they_reach_a_result(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results(
        {
            "test_case_id": "TC-001",
            "source_error": (
                "Login failed. DRIVER={ODBC Driver 18};SERVER=prod-sql-01;"
                "UID=svc_recon;PWD=Sup3rSecret!;"
            ),
        }
    )

    summary, _ = _run(schema, path, book)

    remarks = _by_id(summary)["TC-001"].remarks
    assert "Sup3rSecret!" not in remarks
    assert "svc_recon" not in remarks
    assert "REDACTED" in remarks


def test_comparison_error_is_attributed_to_the_comparison(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 100})

    summary, _ = _run(schema, path, book)

    result = _by_id(summary)["TC-001"]
    assert result.status is ExecutionStatus.ERROR
    assert result.error_side is ErrorSide.COMPARISON
    assert result.error_code == "COMPARISON_FAILED"
    assert "NULL" in result.remarks


def test_unsafe_sql_is_rejected_before_any_executor_is_created(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001", source_sql="DELETE FROM PAYMENTS")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    summary, factory = _run(schema, path, book)

    result = _by_id(summary)["TC-001"]
    assert result.status is ExecutionStatus.ERROR
    assert result.error_side is ErrorSide.SOURCE
    assert result.error_code == "SQL_REJECTED"
    assert factory.created == (), "no connection is opened for a rejected query"


def test_unsafe_target_sql_is_attributed_to_the_target(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001", target_sql="SELECT 1; DROP TABLE t")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    summary, _ = _run(schema, path, book)

    assert _by_id(summary)["TC-001"].error_side is ErrorSide.TARGET


def test_invalid_row_becomes_an_error_result_with_the_workbook_side(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-BAD", source_type="postgres"), case_row("TC-OK")])
    book = make_results({"test_case_id": "TC-OK", "source_result": 1, "target_result": 1})

    summary, _ = _run(schema, path, book)

    bad = _by_id(summary)["TC-BAD"]
    assert bad.status is ExecutionStatus.ERROR
    assert bad.error_side is ErrorSide.WORKBOOK
    assert bad.error_code == "WORKBOOK_ROW_INVALID"
    assert _by_id(summary)["TC-OK"].status is ExecutionStatus.PASS


def test_disabled_cases_are_recorded_as_skipped_and_never_executed(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001"), case_row("TC-002", enabled=False)])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    summary, factory = _run(schema, path, book)

    assert _by_id(summary)["TC-002"].status is ExecutionStatus.SKIPPED
    assert all(e.test_case_id != "TC-002" for e in factory.created)
    assert summary.skipped == 1


def test_executors_are_closed_after_a_successful_run(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    _, factory = _run(schema, path, book)

    assert factory.created
    assert all(e.closed for e in factory.created)


def test_executors_are_closed_after_a_failing_run(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_error": "boom"})

    _, factory = _run(schema, path, book)

    assert factory.created
    assert all(e.closed for e in factory.created)


def test_executors_are_closed_when_writing_fails(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any, tmp_path: Path
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})
    factory = FakeExecutorFactory(book)

    with pytest.raises(WorkbookError):
        ReconciliationRunner(schema, factory).run(path, RunOptions(output_dir=tmp_path / "missing"))

    assert all(e.closed for e in factory.created)


def test_case_selection_runs_only_the_named_case(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001"), case_row("TC-PAY-008"), case_row("TC-003")])
    book = make_results(
        {"test_case_id": "TC-PAY-008", "source_result": 1, "target_result": 1},
        on_unknown_case="error",
    )

    summary, factory = _run(schema, path, book, case_ids=("TC-PAY-008",))

    results = _by_id(summary)
    assert results["TC-PAY-008"].status is ExecutionStatus.PASS
    assert results["TC-001"].status is ExecutionStatus.SKIPPED
    assert results["TC-001"].remarks == "Not selected by --case"
    assert {e.test_case_id for e in factory.created} == {"TC-PAY-008"}


def test_case_selection_is_case_insensitive(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-PAY-008")])
    book = make_results({"test_case_id": "TC-PAY-008", "source_result": 1, "target_result": 1})

    summary, _ = _run(schema, path, book, case_ids=("tc-pay-008",))

    assert summary.passed == 1


def test_unknown_case_selection_is_an_error(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    with pytest.raises(WorkbookError, match="No enabled test case matches --case TC-NOPE"):
        _run(schema, path, book, case_ids=("TC-NOPE",))


def test_limit_runs_a_pilot_subset(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row(f"TC-{i:03d}") for i in range(1, 6)])
    book = make_results(
        *[
            {"test_case_id": f"TC-{i:03d}", "source_result": 1, "target_result": 1}
            for i in range(1, 6)
        ]
    )

    summary, factory = _run(schema, path, book, limit=2)

    assert summary.passed == 2
    assert summary.skipped == 3
    assert _by_id(summary)["TC-003"].remarks == "Beyond --limit"
    assert {e.test_case_id for e in factory.created} == {"TC-001", "TC-002"}


def test_limit_of_zero_executes_nothing(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    summary, factory = _run(schema, path, book, limit=0)

    assert summary.executed == 0
    assert factory.created == ()


def test_negative_limit_is_rejected(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    with pytest.raises(WorkbookError, match="--limit must be 0 or greater"):
        _run(schema, path, book, limit=-1)


def _five_cases(make_workbook: Any, case_row: Any, make_results: Any) -> tuple[Path, Any]:
    path = make_workbook([case_row(f"TC-{i:03d}") for i in range(1, 6)])
    book = make_results(
        *[
            {"test_case_id": f"TC-{i:03d}", "source_result": 1, "target_result": 1}
            for i in range(1, 6)
        ]
    )
    return path, book


def test_from_and_to_case_run_a_slice_from_the_middle(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path, book = _five_cases(make_workbook, case_row, make_results)

    summary, factory = _run(schema, path, book, from_case="TC-002", to_case="tc-004")

    assert summary.passed == 3
    assert summary.skipped == 2
    assert _by_id(summary)["TC-005"].remarks == "Outside --from-case/--to-case range"
    assert {e.test_case_id for e in factory.created} == {"TC-002", "TC-003", "TC-004"}


def test_from_case_alone_runs_to_the_end(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path, book = _five_cases(make_workbook, case_row, make_results)

    _, factory = _run(schema, path, book, from_case="TC-004")

    assert {e.test_case_id for e in factory.created} == {"TC-004", "TC-005"}


def test_to_case_alone_runs_from_the_start(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path, book = _five_cases(make_workbook, case_row, make_results)

    _, factory = _run(schema, path, book, to_case="TC-002")

    assert {e.test_case_id for e in factory.created} == {"TC-001", "TC-002"}


def test_range_is_applied_before_limit(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path, book = _five_cases(make_workbook, case_row, make_results)

    _, factory = _run(schema, path, book, from_case="TC-003", limit=1)

    assert {e.test_case_id for e in factory.created} == {"TC-003"}


def test_disabled_case_can_be_a_range_boundary(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook(
        [case_row("TC-001"), case_row("TC-002", enabled=False), case_row("TC-003")]
    )
    book = make_results(
        {"test_case_id": "TC-001", "source_result": 1, "target_result": 1},
        {"test_case_id": "TC-003", "source_result": 1, "target_result": 1},
    )

    _, factory = _run(schema, path, book, from_case="TC-002")

    assert {e.test_case_id for e in factory.created} == {"TC-003"}


def test_unknown_range_boundary_is_an_error(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path, book = _five_cases(make_workbook, case_row, make_results)

    with pytest.raises(WorkbookError, match="No test case matches --to-case TC-NOPE"):
        _run(schema, path, book, to_case="TC-NOPE")


def test_reversed_range_is_an_error(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path, book = _five_cases(make_workbook, case_row, make_results)

    with pytest.raises(WorkbookError, match="comes after"):
        _run(schema, path, book, from_case="TC-004", to_case="TC-002")


def test_range_cannot_be_combined_with_case(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path, book = _five_cases(make_workbook, case_row, make_results)

    with pytest.raises(WorkbookError, match="cannot be combined"):
        _run(schema, path, book, case_ids=("TC-001",), from_case="TC-002")


def test_fail_fast_stops_after_the_first_problem(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row(f"TC-{i:03d}") for i in range(1, 5)])
    book = make_results(
        {"test_case_id": "TC-001", "source_result": 1, "target_result": 1},
        {"test_case_id": "TC-002", "source_result": 1, "target_result": 9},
        {"test_case_id": "TC-003", "source_result": 1, "target_result": 1},
        {"test_case_id": "TC-004", "source_result": 1, "target_result": 1},
    )

    summary, factory = _run(schema, path, book, fail_fast=True)

    results = _by_id(summary)
    assert results["TC-001"].status is ExecutionStatus.PASS
    assert results["TC-002"].status is ExecutionStatus.FAIL
    assert results["TC-003"].status is ExecutionStatus.SKIPPED
    assert results["TC-003"].remarks == "Stopped by --fail-fast"
    assert results["TC-004"].status is ExecutionStatus.SKIPPED
    assert {e.test_case_id for e in factory.created} == {"TC-001", "TC-002"}


def test_without_fail_fast_every_case_runs(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row(f"TC-{i:03d}") for i in range(1, 4)])
    book = make_results(
        {"test_case_id": "TC-001", "source_result": 1, "target_result": 9},
        {"test_case_id": "TC-002", "source_error": "boom"},
        {"test_case_id": "TC-003", "source_result": 1, "target_result": 1},
    )

    summary, _ = _run(schema, path, book)

    assert (summary.passed, summary.failed, summary.errored) == (1, 1, 1)
    assert summary.executed == 3


def test_results_are_written_to_a_new_timestamped_workbook(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = Path(make_workbook([case_row("TC-001")]))
    before = path.read_bytes()
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    summary, _ = _run(schema, path, book)

    assert summary.output_path is not None
    assert summary.output_path.exists()
    assert summary.output_path != path
    assert "_results_" in summary.output_path.name
    assert path.read_bytes() == before, "the input workbook must never change"


def test_no_write_option_skips_the_output_workbook(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any, tmp_path: Path
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    summary, _ = _run(schema, path, book, write_output=False)

    assert summary.output_path is None
    assert list(tmp_path.glob("*_results_*.xlsx")) == []


def test_durations_and_run_id_are_recorded(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "target_result": 1})

    summary, _ = _run(schema, path, book, run_id="fixedrunid001")

    result = _by_id(summary)["TC-001"]
    assert result.run_id == "fixedrunid001"
    assert result.duration_ms >= 0
    assert result.executed_at.tzinfo is not None
    assert result.executed_at.microsecond == 0


def test_results_are_ordered_by_workbook_row(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    path = make_workbook(
        [case_row("TC-001"), case_row("TC-002", enabled=False), case_row("TC-003")]
    )
    book = make_results(
        {"test_case_id": "TC-001", "source_result": 1, "target_result": 1},
        {"test_case_id": "TC-003", "source_result": 1, "target_result": 1},
    )

    summary, _ = _run(schema, path, book)

    assert [r.row_number for r in summary.results] == [2, 3, 4]


def test_multi_row_result_is_rejected_as_an_error(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any
) -> None:
    """A query returning many rows must stop the case, never leak a result set."""
    path = make_workbook([case_row("TC-001")])
    book = make_results({"test_case_id": "TC-001", "source_result": 1, "source_row_count": 25})

    summary, _ = _run(schema, path, book)

    result = _by_id(summary)["TC-001"]
    assert result.status is ExecutionStatus.ERROR
    assert result.error_side is ErrorSide.SOURCE
    assert "25 rows" in result.remarks


# -- progressive result workbook ---------------------------------------------


def _statuses(output: Path, schema: Any) -> dict[int, Any]:
    """Status cell of every data row in a saved result workbook, by row number."""
    workbook = load_workbook(output)
    try:
        sheet = workbook[schema.sheet_name]
        wanted = schema.fields["status"].header
        column = next(c.column for c in sheet[schema.header_row] if c.value == wanted)
        return {
            row: sheet.cell(row=row, column=column).value
            for row in range(schema.first_data_row, sheet.max_row + 1)
        }
    finally:
        workbook.close()


class _WatchingFactory(FakeExecutorFactory):
    """Calls ``on_case`` just before a case's source query is handed out."""

    def __init__(self, book: Any, on_case: Callable[[str], None]) -> None:
        super().__init__(book)
        self._on_case = on_case

    def get_executor(self, **kwargs: Any) -> Any:
        if kwargs["side"] is QuerySide.SOURCE:
            self._on_case(kwargs["test_case_id"])
        return super().get_executor(**kwargs)


def _three_cases(make_workbook: Any, case_row: Any, make_results: Any) -> tuple[Path, Any]:
    path = make_workbook([case_row(f"TC-{i:03d}") for i in range(1, 4)])
    book = make_results(
        *[
            {"test_case_id": f"TC-{i:03d}", "source_result": 1, "target_result": 1}
            for i in range(1, 4)
        ]
    )
    return path, book


def test_result_workbook_is_updated_after_every_case(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any, tmp_path: Path
) -> None:
    path, book = _three_cases(make_workbook, case_row, make_results)
    seen: dict[str, dict[int, Any]] = {}
    notes: list[str] = []

    def snapshot(test_case_id: str) -> None:
        (output,) = tmp_path.glob("*_results_*.xlsx")
        seen[test_case_id] = _statuses(output, schema)

    runner = ReconciliationRunner(schema, _WatchingFactory(book, snapshot))
    summary = runner.run(path, RunOptions(notify=notes.append))

    assert seen["TC-001"] == {2: None, 3: None, 4: None}, "created before the first query"
    assert seen["TC-002"] == {2: "PASS", 3: None, 4: None}
    assert seen["TC-003"] == {2: "PASS", 3: "PASS", 4: None}
    assert summary.output_path is not None
    assert _statuses(summary.output_path, schema) == {2: "PASS", 3: "PASS", 4: "PASS"}
    assert any(str(summary.output_path) in note for note in notes)


def test_interrupted_run_keeps_the_cases_that_finished(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any, tmp_path: Path
) -> None:
    path, book = _three_cases(make_workbook, case_row, make_results)

    def interrupt(test_case_id: str) -> None:
        if test_case_id == "TC-003":
            raise KeyboardInterrupt

    factory = _WatchingFactory(book, interrupt)
    with pytest.raises(KeyboardInterrupt):
        ReconciliationRunner(schema, factory).run(path, RunOptions())

    (output,) = tmp_path.glob("*_results_*.xlsx")
    assert _statuses(output, schema) == {2: "PASS", 3: "PASS", 4: None}
    assert list(tmp_path.glob("*.tmp")) == [], "no temporary file may be left behind"
    assert all(executor.closed for executor in factory.created)


def test_a_locked_result_workbook_does_not_stop_the_run(
    make_workbook: Any,
    case_row: Any,
    schema: Any,
    make_results: Any,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Excel on Windows locks an open workbook, so replacing it fails."""
    path, book = _three_cases(make_workbook, case_row, make_results)
    real_replace = os.replace

    def locked_once_created(source: Any, destination: Any) -> None:
        if Path(destination).exists():
            Path(source).unlink()
            raise PermissionError(13, "Permission denied")
        real_replace(source, destination)

    monkeypatch.setattr("migration_reconciliation.workbook.writer.os.replace", locked_once_created)
    notes: list[str] = []

    summary, _ = _run(schema, path, book, notify=notes.append)

    assert summary.passed == 3
    assert summary.output_path is not None
    assert summary.output_path.stem.endswith("_2")
    assert _statuses(summary.output_path, schema) == {2: "PASS", 3: "PASS", 4: "PASS"}
    assert sum("Is it open in Excel?" in note for note in notes) == 1
    assert list(tmp_path.glob("*.tmp")) == []


def test_missing_output_directory_fails_before_any_query(
    make_workbook: Any, case_row: Any, schema: Any, make_results: Any, tmp_path: Path
) -> None:
    path, book = _three_cases(make_workbook, case_row, make_results)
    factory = FakeExecutorFactory(book)

    with pytest.raises(WorkbookError, match="Output directory does not exist"):
        ReconciliationRunner(schema, factory).run(path, RunOptions(output_dir=tmp_path / "missing"))

    assert factory.created == ()
