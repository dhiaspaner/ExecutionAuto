"""Writing results into a new copy of the workbook."""

from __future__ import annotations

import hashlib
import re
from dataclasses import replace
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from migration_reconciliation.errors import WorkbookError
from migration_reconciliation.models import (
    ErrorSide,
    ExecutionResult,
    ExecutionStatus,
    WorkbookSchema,
)
from migration_reconciliation.workbook.schema import parse_schema
from migration_reconciliation.workbook.writer import build_output_path, write_results

TIMESTAMP = datetime(2026, 9, 12, 9, 30, 0, tzinfo=UTC)
RUN_ID = "abc123def456"


def _result(**overrides: Any) -> ExecutionResult:
    base = ExecutionResult(
        test_case_id="TC-001",
        row_number=2,
        status=ExecutionStatus.PASS,
        executed_at=TIMESTAMP,
        run_id=RUN_ID,
        duration_ms=42,
        source_result=1500,
        target_result=1500,
        variance=Decimal(0),
        remarks="Source and target numeric results are equal.",
    )
    return replace(base, **overrides)


def _cells(path: Path, schema: WorkbookSchema, row: int = 2) -> dict[str, Any]:
    workbook = load_workbook(path)
    sheet = workbook[schema.sheet_name]
    headers = {cell.value: cell.column for cell in sheet[schema.header_row] if cell.value}
    values = {name: sheet.cell(row=row, column=col).value for name, col in headers.items()}
    workbook.close()
    return values


def test_writes_every_result_field(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row()])

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    cells = _cells(output, schema)
    assert cells["Source Results"] == 1500
    assert cells["Target Results"] == 1500
    assert cells["Variance"] == 0
    assert cells["Status"] == "PASS"
    assert cells["Remarks"] == "Source and target numeric results are equal."
    assert cells["Executed At"] == datetime(2026, 9, 12, 9, 30, 0)
    assert cells["Run ID"] == RUN_ID
    assert cells["Duration (ms)"] == 42
    assert cells["Error Side"] is None


def test_error_side_is_written_for_failures(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row()])

    output = write_results(
        path,
        schema,
        [
            _result(
                status=ExecutionStatus.ERROR,
                error_side=ErrorSide.SOURCE,
                error_code="SOURCE_EXEC_FAILED",
                source_result=None,
                target_result=None,
                variance=None,
                remarks="DatabaseExecutionError: ORA-00942",
            )
        ],
        run_id=RUN_ID,
        timestamp=TIMESTAMP,
    )

    cells = _cells(output, schema)
    assert cells["Status"] == "ERROR"
    assert cells["Error Side"] == "SOURCE"
    assert cells["Error Code"] == "SOURCE_EXEC_FAILED"
    assert cells["Variance"] is None


def test_output_name_is_timestamped(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row()], name="Payment_Domain_V2.xlsx")

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    assert output.name == "Payment_Domain_V2_results_20260912_093000.xlsx"
    assert re.fullmatch(r".+_results_\d{8}_\d{6}\.xlsx", output.name)


def test_the_original_workbook_is_byte_for_byte_unchanged(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = Path(make_workbook([case_row()]))
    digest_before = hashlib.sha256(path.read_bytes()).hexdigest()
    mtime_before = path.stat().st_mtime_ns

    write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    assert hashlib.sha256(path.read_bytes()).hexdigest() == digest_before
    assert path.stat().st_mtime_ns == mtime_before


def test_unrelated_sheets_survive_untouched(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row()])

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Execution Test Cases", "Read Me"]
    assert workbook["Read Me"]["A1"].value == "Migration reconciliation template"
    workbook.close()


def test_input_columns_are_not_disturbed(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row("TC-001", severity="Critical")])

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    cells = _cells(output, schema)
    assert cells["Test Case ID"] == "TC-001"
    assert cells["Source Query"] == "SELECT COUNT(*) FROM PAYMENTS"
    assert cells["Severity"] == "Critical"
    assert cells["Domain"] == "Payments"


def test_rows_without_results_keep_their_result_cells_empty(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = make_workbook([case_row("TC-001"), case_row("TC-002")])

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    assert _cells(output, schema, row=3)["Status"] is None


def test_a_same_second_rerun_gets_its_own_file(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    """Two runs inside one second must not overwrite, and must not lose results."""
    path = make_workbook([case_row()])
    first = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    second = write_results(path, schema, [_result()], run_id="otherrunid99", timestamp=TIMESTAMP)

    assert first.exists() and second.exists()
    assert first != second
    assert second.name.endswith(f"_{'otherrunid99'}.xlsx")


def test_refuses_to_overwrite_when_even_the_run_id_collides(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = make_workbook([case_row()])
    write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)
    write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    with pytest.raises(WorkbookError, match="Refusing to overwrite existing file"):
        write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)


def test_refuses_to_write_over_the_input_workbook(
    make_workbook: Any, case_row: Any, schema_document: dict[str, Any]
) -> None:
    schema_document["workbook"]["output_filename_pattern"] = "{input_stem}{run_id}.xlsx"
    schema = parse_schema(schema_document)
    path = make_workbook([case_row()], workbook_schema=schema, name="cases.xlsx")

    with pytest.raises(WorkbookError, match="Refusing to write results over the input"):
        write_results(path, schema, [_result()], run_id="", timestamp=TIMESTAMP)


def test_write_protection_for_read_only_fields(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    """Defence in depth: a result field marked read-only must never be written."""
    path = make_workbook([case_row()])
    tampered_fields = dict(schema.fields)
    tampered_fields["status"] = replace(tampered_fields["status"], write=False, read=True)
    tampered = replace(schema, fields=tampered_fields)

    with pytest.raises(WorkbookError, match="Refusing to write read-only field 'status'"):
        write_results(path, tampered, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)


def test_missing_required_result_column_is_refused(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema, tmp_path: Path
) -> None:
    path = make_workbook([case_row()])
    workbook = load_workbook(path)
    sheet = workbook[schema.sheet_name]
    for cell in sheet[1]:
        if cell.value == "Status":
            cell.value = None
    stripped = tmp_path / "no_status.xlsx"
    workbook.save(stripped)

    with pytest.raises(WorkbookError, match="missing required result column\\(s\\): Status"):
        write_results(stripped, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)


def test_declared_number_format_is_applied(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row()])

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    workbook = load_workbook(output)
    sheet = workbook[schema.sheet_name]
    column = {c.value: c.column for c in sheet[1]}["Executed At"]
    assert sheet.cell(row=2, column=column).number_format == "yyyy-mm-dd hh:mm:ss"
    workbook.close()


def test_timestamps_are_written_without_a_timezone_excel_cannot_store(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = make_workbook([case_row()])

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    written = _cells(output, schema)["Executed At"]
    assert written.tzinfo is None
    assert written == TIMESTAMP.replace(tzinfo=None), "the UTC instant is preserved"


def test_output_directory_can_be_redirected(
    make_workbook: Any, case_row: Any, schema: Any, tmp_path: Path
) -> None:
    path = make_workbook([case_row()])
    destination = tmp_path / "results"
    destination.mkdir()

    output = write_results(
        path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP, output_dir=destination
    )

    assert output.parent == destination.resolve()


def test_missing_output_directory_is_reported(
    make_workbook: Any, case_row: Any, schema: Any, tmp_path: Path
) -> None:
    path = make_workbook([case_row()])

    with pytest.raises(WorkbookError, match="Output directory does not exist"):
        write_results(
            path,
            schema,
            [_result()],
            run_id=RUN_ID,
            timestamp=TIMESTAMP,
            output_dir=tmp_path / "absent",
        )


def test_other_sheets_can_be_dropped_when_the_schema_says_so(
    make_workbook: Any, case_row: Any, schema_document: dict[str, Any]
) -> None:
    schema_document["workbook"]["preserve_other_sheets"] = False
    schema = parse_schema(schema_document)
    path = make_workbook([case_row()], workbook_schema=schema)

    output = write_results(path, schema, [_result()], run_id=RUN_ID, timestamp=TIMESTAMP)

    workbook = load_workbook(output)
    assert workbook.sheetnames == ["Execution Test Cases"]
    workbook.close()


def test_build_output_path_scrubs_unsafe_substituted_values(
    schema: WorkbookSchema, tmp_path: Path
) -> None:
    output = build_output_path(
        tmp_path / "a/b/../report.xlsx",
        schema,
        timestamp=TIMESTAMP,
        run_id="../../escape",
    )

    assert output.parent == (tmp_path / "a/b/..").resolve()
    assert "/" not in output.name and "\\" not in output.name
