"""Reading test cases from a workbook by header name."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from migration_reconciliation.errors import WorkbookError
from migration_reconciliation.models import (
    ComparisonRule,
    DatabaseType,
    ExecutionScope,
    WorkbookSchema,
)
from migration_reconciliation.workbook.reader import read_workbook, validate_workbook
from migration_reconciliation.workbook.schema import load_schema, parse_schema
from tests.conftest import EXAMPLE_TEMPLATE_PATH, GENERIC_SCHEMA_PATH


def test_reads_the_shipped_template() -> None:
    schema = load_schema(GENERIC_SCHEMA_PATH)

    read = read_workbook(EXAMPLE_TEMPLATE_PATH, schema)

    assert [c.test_case_id for c in read.test_cases][:3] == [
        "TC-PAY-001",
        "TC-PAY-002",
        "TC-PAY-003",
    ]
    assert read.invalid == []


def test_parses_every_semantic_field(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row("TC-001", tolerance="1.5", timeout_seconds=60)])

    case = read_workbook(path, schema).test_cases[0]

    assert case.test_case_id == "TC-001"
    assert case.row_number == 2
    assert case.source_type is DatabaseType.ORACLE
    assert case.source_connection == "LEGACY_ORACLE"
    assert case.target_connection == "MIGRATED_SQLSERVER"
    assert case.source_sql == "SELECT COUNT(*) FROM PAYMENTS"
    assert case.target_sql == "SELECT COUNT(*) FROM dbo.Payments"
    assert case.comparison_rule is ComparisonRule.EQUAL
    assert case.tolerance == Decimal("1.5")
    assert case.timeout_seconds == 60
    assert case.extras["domain"] == "Payments"


def test_missing_sheet_names_the_available_sheets(
    make_workbook: Any, case_row: Any, schema_document: dict[str, Any]
) -> None:
    path = make_workbook([case_row()])
    schema_document["workbook"]["test_case_sheet"] = "Not There"
    schema = parse_schema(schema_document)

    with pytest.raises(WorkbookError, match="Available sheets: Execution Test Cases, Read Me"):
        read_workbook(path, schema)


def test_missing_required_column_is_reported(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema, tmp_path: Path
) -> None:
    path = make_workbook([case_row()])
    workbook = load_workbook(path)
    sheet = workbook[schema.sheet_name]
    for cell in sheet[1]:
        if cell.value == "Source Query":
            cell.value = None
    stripped = tmp_path / "stripped.xlsx"
    workbook.save(stripped)

    with pytest.raises(WorkbookError, match="missing required column\\(s\\): Source Query"):
        read_workbook(stripped, schema)


def test_missing_optional_column_is_tolerated(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema, tmp_path: Path
) -> None:
    path = make_workbook([case_row()])
    workbook = load_workbook(path)
    sheet = workbook[schema.sheet_name]
    for cell in sheet[1]:
        if cell.value == "Severity":
            cell.value = None
    trimmed = tmp_path / "trimmed.xlsx"
    workbook.save(trimmed)

    read = read_workbook(trimmed, schema)

    assert len(read.test_cases) == 1


def test_duplicate_test_case_ids_stop_the_run(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = make_workbook([case_row("TC-001"), case_row("TC-002"), case_row("TC-001")])

    with pytest.raises(WorkbookError, match="duplicate test case id\\(s\\): 'TC-001'"):
        read_workbook(path, schema)


def test_duplicate_headers_in_the_sheet_are_rejected(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema, tmp_path: Path
) -> None:
    path = make_workbook([case_row()])
    workbook = load_workbook(path)
    sheet = workbook[schema.sheet_name]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(row=1, column=len(headers) + 1, value="Status")
    duplicated = tmp_path / "dupe.xlsx"
    workbook.save(duplicated)

    with pytest.raises(WorkbookError, match="duplicate header 'Status'"):
        read_workbook(duplicated, schema)


def test_disabled_cases_are_skipped_not_executed(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = make_workbook([case_row("TC-001", enabled=True), case_row("TC-002", enabled=False)])

    read = read_workbook(path, schema)

    assert [c.test_case_id for c in read.test_cases] == ["TC-001"]
    assert [s.test_case_id for s in read.skipped] == ["TC-002"]
    assert read.skipped[0].reason == "Disabled in the workbook"


@pytest.mark.parametrize("value", ["FALSE", "false", "No", "n", 0, "0", "disabled"])
def test_various_spellings_of_disabled(
    make_workbook: Any, case_row: Any, schema: Any, value: Any
) -> None:
    path = make_workbook([case_row("TC-001", enabled=value)])

    assert read_workbook(path, schema).test_cases == []


@pytest.mark.parametrize("value", ["TRUE", "yes", "Y", 1, "1", "enabled"])
def test_various_spellings_of_enabled(
    make_workbook: Any, case_row: Any, schema: Any, value: Any
) -> None:
    path = make_workbook([case_row("TC-001", enabled=value)])

    assert len(read_workbook(path, schema).test_cases) == 1


def test_defaults_are_applied_when_optional_cells_are_empty(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    row = case_row("TC-001")
    for optional in ("enabled", "comparison_rule", "tolerance", "timeout_seconds"):
        del row[optional]
    path = make_workbook([row])

    case = read_workbook(path, schema).test_cases[0]

    assert case.enabled is True
    assert case.comparison_rule is ComparisonRule.EQUAL
    assert case.tolerance == Decimal(0)
    assert case.timeout_seconds == 120


def test_alternate_header_names_need_no_code_change(
    schema_document: dict[str, Any], case_row: Any, make_workbook: Any
) -> None:
    """The same semantic fields, bound to a completely different template."""
    schema_document["profile_name"] = "fines-domain"
    schema_document["workbook"]["test_case_sheet"] = "Fines Recon"
    schema_document["workbook"]["header_row"] = 3
    schema_document["workbook"]["first_data_row"] = 5
    schema_document["fields"]["test_case_id"]["header"] = "Scenario Ref"
    schema_document["fields"]["source_sql"]["header"] = "Legacy Query"
    schema_document["fields"]["target_sql"]["header"] = "New Query"
    schema_document["fields"]["source_type"]["header"] = "Legacy Engine"
    alternate = parse_schema(schema_document)
    path = make_workbook([case_row("FINE-01")], workbook_schema=alternate, name="fines.xlsx")

    read = read_workbook(path, alternate)

    assert [c.test_case_id for c in read.test_cases] == ["FINE-01"]
    assert read.test_cases[0].row_number == 5
    assert read.test_cases[0].source_sql == "SELECT COUNT(*) FROM PAYMENTS"


def test_header_matching_ignores_case_and_surrounding_whitespace(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema, tmp_path: Path
) -> None:
    path = make_workbook([case_row()])
    workbook = load_workbook(path)
    sheet = workbook[schema.sheet_name]
    for cell in sheet[1]:
        if cell.value == "Status":
            cell.value = "  status  "
    retitled = tmp_path / "retitled.xlsx"
    workbook.save(retitled)

    assert len(read_workbook(retitled, schema).test_cases) == 1


@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"source_type": "postgres"}, "unsupported value 'postgres'"),
        ({"comparison_rule": "close_enough"}, "unsupported value 'close_enough'"),
        ({"tolerance": "abc"}, "must be a number"),
        ({"timeout_seconds": "soon"}, "must be a whole number"),
        ({"timeout_seconds": 1.5}, "must be a whole number"),
        ({"enabled": "maybe"}, "must be TRUE or FALSE"),
        ({"source_sql": None}, "is required but empty"),
        ({"source_connection": "   "}, "is required but empty"),
    ],
)
def test_bad_row_becomes_an_invalid_row_not_a_crash(
    make_workbook: Any, case_row: Any, schema: Any, overrides: dict[str, Any], message: str
) -> None:
    path = make_workbook([case_row("TC-BAD", **overrides), case_row("TC-GOOD")])

    read = read_workbook(path, schema)

    assert [c.test_case_id for c in read.test_cases] == ["TC-GOOD"], "good rows still run"
    assert len(read.invalid) == 1
    assert read.invalid[0].test_case_id == "TC-BAD"
    assert message in read.invalid[0].message


def test_row_without_an_id_is_invalid(make_workbook: Any, case_row: Any, schema: Any) -> None:
    path = make_workbook([case_row(""), case_row("TC-GOOD")])

    read = read_workbook(path, schema)

    assert len(read.invalid) == 1
    assert "Test case id is empty" in read.invalid[0].message


def test_blank_rows_are_ignored(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema, tmp_path: Path
) -> None:
    path = make_workbook([case_row("TC-001"), case_row("TC-002")])
    workbook = load_workbook(path)
    sheet = workbook[schema.sheet_name]
    sheet.insert_rows(3)
    spaced = tmp_path / "spaced.xlsx"
    workbook.save(spaced)

    read = read_workbook(spaced, schema)

    assert [c.test_case_id for c in read.test_cases] == ["TC-001", "TC-002"]
    assert read.invalid == []


def test_enum_values_are_matched_case_insensitively(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = make_workbook([case_row("TC-001", source_type="Oracle", comparison_rule="EQUAL")])

    case = read_workbook(path, schema).test_cases[0]

    assert case.source_type is DatabaseType.ORACLE
    assert case.comparison_rule is ComparisonRule.EQUAL


def test_missing_workbook_reports_a_clear_error(schema: Any, tmp_path: Path) -> None:
    with pytest.raises(WorkbookError, match="Workbook not found"):
        read_workbook(tmp_path / "absent.xlsx", schema)


def test_validate_workbook_does_not_modify_the_file(
    make_workbook: Any, case_row: Any, schema: Any
) -> None:
    path = make_workbook([case_row()])
    before = Path(path).read_bytes()

    validate_workbook(path, schema)

    assert Path(path).read_bytes() == before


# ---------------------------------------------------------------------------
# One-sided tests: a scope that names a single side is not asked for the other.
# ---------------------------------------------------------------------------


def test_target_only_row_needs_no_source_connection(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    """A TARGET_ONLY row leaves the whole source side empty and still reads."""
    path = make_workbook(
        [
            case_row(
                "TC-TGT-001",
                execution_scope="TARGET_ONLY",
                source_connection="",
                source_sql="",
                comparison_rule="expected_zero",
            )
        ]
    )

    read = read_workbook(path, schema)

    assert read.invalid == []
    assert len(read.test_cases) == 1
    assert read.test_cases[0].execution_scope is ExecutionScope.TARGET_ONLY
    assert read.test_cases[0].source_sql == ""


def test_source_only_row_needs_no_target_connection(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    path = make_workbook(
        [
            case_row(
                "TC-SRC-001",
                execution_scope="SOURCE_ONLY",
                target_connection="",
                target_sql="",
                comparison_rule="expected_zero",
            )
        ]
    )

    read = read_workbook(path, schema)

    assert read.invalid == []
    assert read.test_cases[0].execution_scope is ExecutionScope.SOURCE_ONLY
    assert read.test_cases[0].target_sql == ""


def test_a_populated_unused_side_is_rejected(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    """The row and its scope disagree, so the row is reported rather than guessed."""
    path = make_workbook(
        [
            case_row(
                "TC-TGT-002",
                execution_scope="TARGET_ONLY",
                source_connection="",
                source_sql="SELECT COUNT(*) FROM PAYMENTS",
            )
        ]
    )

    read = read_workbook(path, schema)

    assert read.test_cases == []
    assert len(read.invalid) == 1
    assert "must be empty when Execution_Scope is TARGET_ONLY" in read.invalid[0].message


def test_both_sides_are_still_required_by_default(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    """Omitting the scope keeps the two-sided behaviour, so nothing loosens silently."""
    path = make_workbook([case_row("TC-BOTH-001", source_sql="")])

    read = read_workbook(path, schema)

    assert read.test_cases == []
    assert len(read.invalid) == 1
    assert "is required but empty" in read.invalid[0].message


# ---------------------------------------------------------------------------
# A disabled row says why, in the workbook's own words.
# ---------------------------------------------------------------------------


def test_a_disabled_row_quotes_the_reason_from_the_workbook(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    path = make_workbook(
        [
            case_row(
                "TC-OFF",
                enabled=False,
                test_name="No mapped Arabic-language attribute exists for this target table.",
            )
        ]
    )

    read = read_workbook(path, schema)

    assert len(read.skipped) == 1
    assert read.skipped[0].reason == (
        "Disabled in the workbook: No mapped Arabic-language attribute exists "
        "for this target table."
    )


def test_a_disabled_row_without_a_name_still_says_it_is_disabled(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    path = make_workbook([case_row("TC-OFF", enabled=False)])

    read = read_workbook(path, schema)

    assert read.skipped[0].reason == "Disabled in the workbook"


def test_a_very_long_reason_is_truncated(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    """One enormous cell must not crowd out the rest of a result sheet."""
    path = make_workbook([case_row("TC-OFF", enabled=False, test_name="x" * 500)])

    read = read_workbook(path, schema)

    assert read.skipped[0].reason.endswith("...")
    assert len(read.skipped[0].reason) < 400


def test_an_enabled_row_is_unaffected_by_its_name(
    make_workbook: Any, case_row: Any, schema: WorkbookSchema
) -> None:
    path = make_workbook([case_row("TC-ON", test_name="Row counts match")])

    read = read_workbook(path, schema)

    assert read.skipped == []
    assert len(read.test_cases) == 1
