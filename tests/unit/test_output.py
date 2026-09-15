"""Writing results: what is written, what is preserved, and how it is saved."""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

import migration_reconciliation.workbook.output as output_module
from migration_reconciliation.errors import SqlSyntaxError, WorkbookError
from migration_reconciliation.execution.engine import execute_plan
from migration_reconciliation.execution.plan import ExecutionPlan, build_plan
from migration_reconciliation.workbook.output import ResultWriter, resolve_directory_env
from migration_reconciliation.workbook.run_control import OutputMode
from tests.scripted import ScriptedExecutors


def cells(path: Path, sheet_name: str = "Test Cases") -> dict[str, dict[str, Any]]:
    """The sheet as ``{Test_ID: {header: value}}``."""
    workbook = load_workbook(path, data_only=True)
    try:
        sheet = workbook[sheet_name]
        headers = [cell.value for cell in sheet[1]]
        rows: dict[str, dict[str, Any]] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = dict(zip(headers, row, strict=False))
            if values.get("Test_ID"):
                rows[str(values["Test_ID"])] = values
        return rows
    finally:
        workbook.close()


def history_column(workbook: Any, header: str) -> int:
    """Index of a ``Run History`` column, so a test never counts columns by hand."""
    headers = [cell.value for cell in workbook["Run History"][1]]
    return headers.index(header)


def run_and_write(path: Path, executors: ScriptedExecutors, **kwargs: Any) -> Any:
    return execute_plan(
        build_plan(path), executors=executors, run_id="run-1", write_output=True, **kwargs
    )


# -- what lands in the workbook ----------------------------------------------


def test_a_run_writes_every_output_column_it_finds(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])

    report = run_and_write(path, ScriptedExecutors(results={"source": 100, "target": 97}))

    written = cells(report.output_path)["TC-001"]
    assert written["Status"] == "FAIL"
    assert written["Source_Result"] == 100
    assert written["Target_Result"] == 97
    assert written["Actual_Value"] == 97
    assert written["Variance"] == -3
    assert written["Run_ID"] == "run-1"
    assert written["Error_Code"] == "VALUE_MISMATCH"
    assert written["Observation"]
    assert written["Executed_At_UTC"] is not None
    assert written["Source_Duration_ms"] is not None


def test_definition_columns_are_never_modified(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    before = cells(path)["TC-001"]

    report = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 2}))

    after = cells(report.output_path)["TC-001"]
    for column in (
        "Test_ID",
        "Enabled",
        "Domain",
        "Execution_Scope",
        "Comparison_Type",
        "Result_Type",
        "Source_SQL",
        "Target_SQL",
        "Severity",
    ):
        assert after[column] == before[column], column


def test_the_input_workbook_is_left_exactly_as_it_was(
    make_recon_workbook: Any, test_row: Any
) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    original = path.read_bytes()

    report = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 1}))

    assert path.read_bytes() == original
    assert report.output_path != path
    assert cells(path)["TC-001"]["Status"] is None


def test_unrelated_sheets_survive_a_run(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])

    report = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 1}))

    workbook = load_workbook(report.output_path)
    try:
        assert "Executor Contract" in workbook.sheetnames
        assert "Conversion Notes" in workbook.sheetnames
        assert "Connections" in workbook.sheetnames
        assert workbook["Executor Contract"]["A1"].value.startswith("Executor contract")
    finally:
        workbook.close()


def test_a_disabled_row_gets_the_disabled_status(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-OFF", Enabled="No")])

    report = execute_plan(
        build_plan(path),
        executors=ScriptedExecutors(results={}),
        run_id="run-1",
        write_output=True,
    )

    assert cells(report.output_path)["TC-OFF"]["Status"] == "DISABLED"


def test_stale_results_are_cleared_at_the_start_of_a_real_run(
    make_recon_workbook: Any, test_row: Any
) -> None:
    """Last week's PASS must not survive on a row that errors today."""
    path = make_recon_workbook([test_row("TC-001")])
    workbook = load_workbook(path)
    sheet = workbook["Test Cases"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(row=2, column=headers["Status"], value="PASS")
    sheet.cell(row=2, column=headers["Observation"], value="all good last time")
    workbook.save(path)
    workbook.close()

    report = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 2}))

    written = cells(report.output_path)["TC-001"]
    assert written["Status"] == "FAIL"
    assert "all good last time" not in str(written["Observation"])


# -- run history -------------------------------------------------------------


def test_one_row_is_appended_to_run_history(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])

    report = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 1}))

    workbook = load_workbook(report.output_path, data_only=True)
    try:
        history = workbook["Run History"]
        headers = [cell.value for cell in history[1]]
        rows = [
            dict(zip(headers, values, strict=False))
            for values in history.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in values)
        ]
    finally:
        workbook.close()

    assert len(rows) == 1
    entry = rows[0]
    assert entry["Run_ID"] == "run-1"
    assert entry["Overall_Status"] == "PASS"
    assert entry["Enabled_Tests"] == 1
    assert entry["Passed"] == 1
    assert entry["Workbook_Name"] == path.name
    assert entry["Output_File"] == report.output_path.name


def test_history_rows_accumulate_rather_than_overwrite(
    make_recon_workbook: Any, test_row: Any
) -> None:
    path = make_recon_workbook([test_row("TC-001")])

    first = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 1}))
    second = execute_plan(
        build_plan(first.output_path),
        executors=ScriptedExecutors(results={"source": 1, "target": 1}),
        run_id="run-2",
        write_output=True,
    )

    workbook = load_workbook(second.output_path, data_only=True)
    try:
        history = workbook["Run History"]
        ids = [row[0] for row in history.iter_rows(min_row=2, values_only=True) if row[0]]
    finally:
        workbook.close()

    assert ids == ["run-1", "run-2"]


# -- how it is saved ---------------------------------------------------------


def test_nothing_is_left_behind_when_a_save_fails(
    make_recon_workbook: Any, test_row: Any, tmp_path: Path
) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    writer = ResultWriter(path, sheet_name="Test Cases", header_row=1)
    try:
        with pytest.raises(WorkbookError, match="does not exist"):
            writer.save(tmp_path / "missing-folder" / "out.xlsx")
    finally:
        writer.close()

    assert not list(tmp_path.glob("*.tmp"))
    assert not list(tmp_path.glob(".*"))


def test_two_runs_in_one_second_do_not_collide(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])

    first = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 1}))
    second = execute_plan(
        build_plan(path),
        executors=ScriptedExecutors(results={"source": 1, "target": 1}),
        run_id="run-2",
        write_output=True,
    )

    assert first.output_path != second.output_path
    assert first.output_path.exists() and second.output_path.exists()


def test_the_output_directory_can_be_chosen(
    make_recon_workbook: Any, test_row: Any, tmp_path: Path
) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    destination = tmp_path / "results"
    destination.mkdir()

    report = run_and_write(
        path, ScriptedExecutors(results={"source": 1, "target": 1}), output_dir=destination
    )

    assert report.output_path.parent == destination.resolve()


def test_an_output_directory_environment_variable_is_read(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    destination = tmp_path / "out"
    destination.mkdir()
    monkeypatch.setenv("RECON_OUT", str(destination))

    assert resolve_directory_env("RECON_OUT") == destination
    assert resolve_directory_env("") is None
    assert resolve_directory_env("RECON_UNSET") is None


def test_an_output_directory_that_is_not_there_is_refused(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("RECON_OUT", str(tmp_path / "nowhere"))

    with pytest.raises(WorkbookError, match="not an existing directory"):
        resolve_directory_env("RECON_OUT")


def test_in_place_mode_replaces_the_input_through_a_temporary_file(
    make_recon_workbook: Any, test_row: Any, run_control_settings: Any
) -> None:
    path = make_recon_workbook(
        [test_row("TC-001")], run_control=run_control_settings(Output_Mode="IN_PLACE")
    )

    report = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 1}))

    assert report.output_path == path.resolve()
    assert cells(path)["TC-001"]["Status"] == "PASS"
    assert not list(path.parent.glob("*.tmp"))


def test_the_writer_refuses_anything_that_is_not_an_output_column(
    make_recon_workbook: Any, test_row: Any
) -> None:
    """Defence in depth: the caller cannot widen what a run may touch."""
    path = make_recon_workbook([test_row("TC-001")])
    writer = ResultWriter(path, sheet_name="Test Cases", header_row=1)
    try:
        with pytest.raises(WorkbookError, match="not an execution-output column"):
            writer.write_row(2, {"Source_SQL": "SELECT 1"})
    finally:
        writer.close()


def test_a_dry_run_writes_nothing_at_all(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    workbook = load_workbook(path)
    sheet = workbook["Test Cases"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(row=2, column=headers["Status"], value="PASS")
    workbook.save(path)
    workbook.close()
    before = path.read_bytes()
    plan = build_plan(path)

    report = execute_plan(
        ExecutionPlan(
            workbook_path=plan.workbook_path,
            sheet=plan.sheet,
            control=plan.control.with_overrides(dry_run=True),
            rules=plan.rules,
        ),
        executors=None,
        run_id="dry",
        write_output=True,
    )

    assert report.output_path is None
    assert path.read_bytes() == before
    assert cells(path)["TC-001"]["Status"] == "PASS"  # not cleared
    assert not list(path.parent.glob("*_results_*.xlsx"))


def test_output_mode_new_file_is_the_default(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])

    assert build_plan(path).control.output_mode is OutputMode.NEW_FILE


def test_no_credential_ever_reaches_the_workbook(make_recon_workbook: Any, test_row: Any) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    from migration_reconciliation.errors import DatabaseExecutionError

    leaky = DatabaseExecutionError("PWD=hunter2;UID=svc_recon;SERVER=sql01")
    report = run_and_write(path, ScriptedExecutors(results={"source": leaky, "target": 1}))

    text = report.output_path.read_bytes()
    assert b"hunter2" not in text


def test_the_run_survives_being_asked_for_a_status_without_a_history_sheet(
    make_recon_workbook: Any, test_row: Any, tmp_path: Path
) -> None:
    """A workbook without Run History gets one rather than losing the record."""
    from migration_reconciliation.workbook.payments_template import build_workbook

    workbook = build_workbook([test_row("TC-001")])
    del workbook["Run History"]
    path = tmp_path / "no-history.xlsx"
    workbook.save(path)
    workbook.close()

    report = run_and_write(path, ScriptedExecutors(results={"source": 1, "target": 1}))

    result = load_workbook(report.output_path, data_only=True)
    try:
        assert "Run History" in result.sheetnames
        assert result["Run History"]["A2"].value == "run-1"
    finally:
        result.close()


def test_the_temporary_file_lives_beside_the_destination(
    make_recon_workbook: Any, test_row: Any, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A cross-device rename is not atomic, so the temp file shares the folder."""
    path = make_recon_workbook([test_row("TC-001")])
    destination = tmp_path / "results"
    destination.mkdir()
    seen: list[str] = []
    original_replace = os.replace

    def watching_replace(src: Any, dst: Any) -> None:
        seen.append(str(Path(src).parent))
        original_replace(src, dst)

    monkeypatch.setattr(output_module.os, "replace", watching_replace)
    writer = ResultWriter(path, sheet_name="Test Cases", header_row=1)
    try:
        writer.save(destination / "out.xlsx")
    finally:
        writer.close()

    assert seen == [str(destination)]


# -- what a stopped validation pass leaves behind -----------------------------


def test_a_gated_run_writes_the_reasons_and_no_results(
    make_recon_workbook: Any, test_row: Any
) -> None:
    broken = test_row("TC-BAD", Source_SQL="SELECT COUNT(*) FROM Paymnets")
    path = make_recon_workbook([test_row("TC-OK"), broken])
    executors = ScriptedExecutors(
        results={"source": 100, "target": 100},
        syntax_failures={
            ("source", "SELECT COUNT(*) FROM Paymnets"): SqlSyntaxError(
                "Invalid object name 'Paymnets'"
            )
        },
    )

    report = run_and_write(path, executors)

    written = cells(report.output_path)
    assert written["TC-BAD"]["Status"] == "SYNTAX ERROR"
    assert written["TC-BAD"]["Error_Code"] == "SYNTAX_ERROR"
    assert "Paymnets" in written["TC-BAD"]["Error_Detail"]
    assert written["TC-OK"]["Status"] == "NOT EXECUTED"
    for test_id in ("TC-BAD", "TC-OK"):
        assert written[test_id]["Source_Result"] is None, test_id
        assert written[test_id]["Target_Result"] is None, test_id
        assert written[test_id]["Variance"] is None, test_id
    assert not executors.executed_anything()


def test_a_gated_run_still_leaves_the_input_workbook_alone(
    make_recon_workbook: Any, test_row: Any
) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    original = path.read_bytes()

    report = run_and_write(
        path,
        ScriptedExecutors(
            results={"source": 1, "target": 1},
            syntax_failures={("source", test_row("TC-001")["Source_SQL"]): SqlSyntaxError("nope")},
        ),
    )

    assert path.read_bytes() == original
    assert report.output_path is not None and report.output_path != path


def test_a_gated_run_appends_one_history_row_saying_nothing_passed(
    make_recon_workbook: Any, test_row: Any
) -> None:
    path = make_recon_workbook([test_row("TC-001")])
    report = run_and_write(
        path,
        ScriptedExecutors(
            results={"source": 1, "target": 1},
            syntax_failures={("source", test_row("TC-001")["Source_SQL"]): SqlSyntaxError("nope")},
        ),
    )

    workbook = load_workbook(report.output_path)
    try:
        history = list(workbook["Run History"].iter_rows(min_row=2, values_only=True))
        status_column = history_column(workbook, "Overall_Status")
    finally:
        workbook.close()

    assert len(history) == 1
    assert history[0][status_column] == "ERROR"
