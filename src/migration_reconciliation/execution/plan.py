"""Building an execution plan from a workbook, before anything is opened.

The plan is everything that can be known without a database: which tests are
enabled, which are misconfigured, which TOML sections are needed, how long a
query may take, and what the observations should say. A dry run is exactly this
much work and no more, which is what makes ``--dry-run`` a real check rather
than a rehearsal.

The workbook is opened read-only here. The writable copy is opened later, by
the writer, so the file being validated is never the file being modified.

Sheets this module reads: ``Test Cases``, ``Run Control``, ``Observation
Rules``, and ``Comparison Types`` for cross-checking only. Sheets it
deliberately ignores: ``Executor Contract``, ``Conversion Notes`` and
``Connections``. They are documentation and a TOML example; a run takes no
instruction, no action and no connection setting from any of them.
"""

from __future__ import annotations

from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from ..errors import WorkbookError
from ..models import ComparisonType
from ..workbook.columns import (
    COMPARISON_TYPES_SHEET,
    DOCUMENTATION_SHEETS,
    OBSERVATION_RULES_SHEET,
    RUN_CONTROL_SHEET,
    TEST_CASES_SHEET,
    normalize_header,
)
from ..workbook.observations import ObservationRules, read_observation_rules
from ..workbook.run_control import RunControl, read_run_control
from ..workbook.sheetio import (
    cell,
    find_sheet,
    locate_header_row,
    map_headers,
    sheet_rows,
    text_of,
)
from ..workbook.testcases import (
    KNOWN_PROFILE_SECTIONS,
    DefinitionProblem,
    TestCaseSheet,
    TestDefinition,
    read_test_cases,
)

__all__ = ["ExecutionPlan", "build_plan"]


@dataclass(frozen=True, slots=True)
class ExecutionPlan:
    """What a run is about to do, decided entirely offline."""

    workbook_path: Path
    sheet: TestCaseSheet
    control: RunControl
    rules: ObservationRules
    warnings: tuple[str, ...] = ()
    ignored_sheets: tuple[str, ...] = ()

    @property
    def executable(self) -> tuple[TestDefinition, ...]:
        """Enabled, valid tests. These are the only ones that reach a database."""
        return self.sheet.enabled

    @property
    def disabled(self) -> tuple[TestDefinition, ...]:
        return self.sheet.disabled

    @property
    def problems(self) -> tuple[DefinitionProblem, ...]:
        return self.sheet.problems

    @property
    def required_sections(self) -> tuple[str, ...]:
        """The TOML sections to open. Nothing outside this tuple is connected to."""
        return self.sheet.required_sections()

    @property
    def enabled_test_count(self) -> int:
        """Executable tests. Disabled rows are not part of any total."""
        return len(self.executable)


def build_plan(
    workbook_path: str | Path,
    *,
    sheet_name: str = TEST_CASES_SHEET,
    known_sections: Sequence[str] = KNOWN_PROFILE_SECTIONS,
) -> ExecutionPlan:
    """Read and validate everything the run needs, touching no database."""
    path = Path(workbook_path)
    if not path.is_file():
        raise WorkbookError(f"Workbook not found: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise WorkbookError(f"Cannot open workbook '{path.name}': {exc}") from None

    try:
        control = read_run_control(workbook, sheet_name=RUN_CONTROL_SHEET)
        sheet = read_test_cases(workbook, sheet_name=sheet_name, known_sections=known_sections)
        rules = read_observation_rules(workbook, sheet_name=OBSERVATION_RULES_SHEET)
        declared = _declared_comparison_types(workbook)
        ignored = tuple(
            name
            for name in workbook.sheetnames
            if normalize_header(name) in {normalize_header(s) for s in DOCUMENTATION_SHEETS}
        )
    finally:
        workbook.close()

    warnings = list(control.warnings)
    if control.defaulted:
        warnings.append(
            f"No '{RUN_CONTROL_SHEET}' sheet was found; default run settings are being used."
        )
    unimplemented = sorted(declared - {c.value for c in ComparisonType})
    if unimplemented:
        warnings.append(
            f"'{COMPARISON_TYPES_SHEET}' documents comparison type(s) this executor does not "
            f"implement: {', '.join(unimplemented)}. A test that asks for one is reported as "
            f"CONFIG ERROR rather than guessed at."
        )
    if sheet.missing_output_columns:
        warnings.append(
            f"Sheet '{sheet.sheet_name}' has no "
            f"{', '.join(sheet.missing_output_columns)} column(s); those outputs are not written."
        )
    if sheet.legacy_columns:
        warnings.append(
            f"Sheet '{sheet.sheet_name}' carries legacy column(s) "
            f"{', '.join(sheet.legacy_columns)}. They are preserved and never executed: "
            f"status, platform and error code decide what happens."
        )
    if ignored:
        warnings.append(
            f"Documentation sheet(s) {', '.join(ignored)} are never parsed for instructions "
            f"or connection settings."
        )

    return ExecutionPlan(
        workbook_path=path,
        sheet=sheet,
        control=control,
        rules=rules,
        warnings=tuple(warnings),
        ignored_sheets=ignored,
    )


def _declared_comparison_types(workbook: Any) -> set[str]:
    """Read ``Comparison Types`` as metadata, purely to cross-check the workbook.

    Nothing on that sheet defines behaviour. It is read so a template that
    documents a comparison Python does not implement is reported now, instead
    of surfacing later as a row-by-row CONFIG ERROR.
    """
    sheet = find_sheet(workbook, COMPARISON_TYPES_SHEET)
    if sheet is None:
        return set()
    rows = sheet_rows(sheet)
    header_row = locate_header_row(
        rows, ("Comparison_Type", "Description"), must_contain="Comparison_Type"
    )
    if header_row is None:
        return set()
    columns = map_headers(rows[header_row - 1])
    column = columns.get(normalize_header("Comparison_Type"))
    if column is None:  # pragma: no cover - guarded by must_contain
        return set()
    # Stop at the first blank row: the sheet often carries other reference
    # lists below the table, and they are not comparison types.
    declared: set[str] = set()
    for row in rows[header_row:]:
        name = text_of(cell(row, column)).upper()
        if not name:
            break
        declared.add(name)
    return declared
