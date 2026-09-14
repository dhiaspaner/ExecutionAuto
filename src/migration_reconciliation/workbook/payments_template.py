"""Generating a reconciliation workbook that matches this executor exactly.

The template is produced from the same constants the reader uses, so the two
cannot drift: if a column is renamed in :mod:`.columns`, the generated workbook
renames with it. That makes this module the readable specification of the
template as well as a way to get one.

The generated ``Connections`` sheet holds a TOML *example*, and the
``Executor Contract`` and ``Conversion Notes`` sheets hold prose. All three are
inert: the executor never reads a connection setting, an instruction or an
action from any sheet. Connections come from the TOML profile and nowhere else.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import ComparisonType, ErrorCode, ExecutionScope, ResultType, TestStatus
from .columns import (
    COMPARISON_TYPES_SHEET,
    DEFINITION_COLUMNS,
    OBSERVATION_RULES_SHEET,
    OUTPUT_COLUMNS,
    RUN_CONTROL_SHEET,
    RUN_HISTORY_COLUMNS,
    RUN_HISTORY_SHEET,
    TEST_CASES_SHEET,
)

__all__ = ["DEFAULT_RUN_CONTROL", "DEMO_TESTS", "build_workbook", "write_workbook"]

_DEFINITION_FILL = PatternFill("solid", fgColor="DDEBF7")
_OUTPUT_FILL = PatternFill("solid", fgColor="FCE4D6")
_HEADER_FONT = Font(bold=True)

TEMPLATE_VERSION = "2.0"

#: The run-control defaults a fresh template ships with.
DEFAULT_RUN_CONTROL: tuple[tuple[str, Any, str], ...] = (
    ("Template_Version", TEMPLATE_VERSION, "Refused by an executor that does not support it."),
    ("Query_Timeout_Seconds", 120, "Applied independently to every source and target query."),
    ("Max_Parallel_Workers", 1, "This executor runs sequentially and caps the value at 1."),
    ("Continue_On_Test_Error", "Yes", "No stops the run at the first ERROR."),
    (
        "Stop_On_Critical_Config_Error",
        "Yes",
        "Yes refuses to open any database while a test is misconfigured.",
    ),
    ("Observation_Max_Length", 300, "Observations are sanitized and truncated to this."),
    ("Error_Detail_Max_Length", 500, "Error_Detail is sanitized and truncated to this."),
    ("Output_Mode", "NEW_FILE", "NEW_FILE keeps the input workbook. IN_PLACE replaces it."),
    ("Output_Directory_Env", "", "Name of an environment variable holding the output folder."),
    ("Evidence_Directory_Env", "", "Reserved. No evidence file is written: results are scalars."),
    ("Dry_Run", "No", "Yes validates everything and runs no reconciliation SQL."),
    (
        "Require_Read_Only_SQL",
        "Yes",
        "Always enforced. No is ignored: read-only validation is a security requirement.",
    ),
)

#: Illustrative tests covering each scope. The SQL is never run by this module.
DEMO_TESTS: tuple[dict[str, Any], ...] = (
    {
        "Test_ID": "TC-PAY-001",
        "Enabled": "Yes",
        "Domain": "Payments",
        "Flow": "Migration",
        "Test_Name": "Posted payment row count matches",
        "Execution_Scope": "SOURCE_TARGET",
        "Comparison_Type": "EQUAL",
        "Result_Type": "INTEGER",
        "Source_Profile_Section": "source",
        "Source_Object": "webservice.dbo.Payments",
        "Source_SQL": "SELECT COUNT(*) FROM webservice.dbo.Payments WHERE Status = 'POSTED'",
        "Target_Profile_Section": "target",
        "Target_Object": "dbo.Payments",
        "Target_SQL": "SELECT COUNT(*) FROM dbo.Payments WHERE Status = 'POSTED'",
        "Severity": "High",
        "Owner": "migration-team",
        "Tags": "counts",
    },
    {
        "Test_ID": "TC-PAY-002",
        "Enabled": "Yes",
        "Domain": "Payments",
        "Flow": "Migration",
        "Test_Name": "Payment totals agree within rounding",
        "Execution_Scope": "SOURCE_TARGET",
        "Comparison_Type": "EQUAL_ABS_TOLERANCE",
        "Result_Type": "NUMBER",
        "Source_Profile_Section": "source",
        "Source_Object": "webservice.dbo.Payments",
        "Source_SQL": "SELECT COALESCE(SUM(Amount), 0) FROM webservice.dbo.Payments",
        "Target_Profile_Section": "target",
        "Target_Object": "dbo.Payments",
        "Target_SQL": "SELECT COALESCE(SUM(Amount), 0) FROM dbo.Payments",
        "Absolute_Tolerance": "0.01",
        "Severity": "High",
        "Owner": "migration-team",
        "Tags": "totals",
    },
    {
        "Test_ID": "TC-PAY-003",
        "Enabled": "Yes",
        "Domain": "Payments",
        "Flow": "Migration",
        "Test_Name": "Fee totals within one percent",
        "Execution_Scope": "SOURCE_TARGET",
        "Comparison_Type": "EQUAL_PCT_TOLERANCE",
        "Result_Type": "NUMBER",
        "Source_Profile_Section": "source",
        "Source_Object": "etables.dbo.Fees",
        "Source_SQL": "SELECT COALESCE(SUM(Fee), 0) FROM etables.dbo.Fees",
        "Target_Profile_Section": "target",
        "Target_Object": "dbo.Fees",
        "Target_SQL": "SELECT COALESCE(SUM(Fee), 0) FROM dbo.Fees",
        "Percentage_Tolerance": "1",
        "Severity": "Medium",
        "Owner": "migration-team",
        "Tags": "totals",
    },
    {
        "Test_ID": "TC-PAY-004",
        "Enabled": "Yes",
        "Domain": "Payments",
        "Flow": "Data quality",
        "Test_Name": "No orphaned payments in the legacy system",
        "Execution_Scope": "SOURCE_ONLY",
        "Comparison_Type": "EXPECTED_ZERO",
        "Result_Type": "INTEGER",
        "Source_Profile_Section": "source",
        "Source_Object": "webservice.dbo.Payments",
        "Source_SQL": (
            "SELECT COUNT(*) FROM webservice.dbo.Payments p "
            "WHERE NOT EXISTS (SELECT 1 FROM webservice.dbo.Customers c "
            "WHERE c.Id = p.CustomerId)"
        ),
        "Severity": "Critical",
        "Owner": "migration-team",
        "Tags": "integrity",
    },
    {
        "Test_ID": "TC-PAY-005",
        "Enabled": "Yes",
        "Domain": "Payments",
        "Flow": "Data quality",
        "Test_Name": "Migrated system has no duplicate references",
        "Execution_Scope": "TARGET_ONLY",
        "Comparison_Type": "EXPECTED_ZERO",
        "Result_Type": "INTEGER",
        "Target_Profile_Section": "target",
        "Target_Object": "dbo.Payments",
        "Target_SQL": (
            "SELECT COUNT(*) FROM (SELECT Reference FROM dbo.Payments "
            "GROUP BY Reference HAVING COUNT(*) > 1) d"
        ),
        "Severity": "Critical",
        "Owner": "migration-team",
        "Tags": "integrity",
    },
    {
        "Test_ID": "TC-PAY-006",
        "Enabled": "Yes",
        "Domain": "Payments",
        "Flow": "Profiling",
        "Test_Name": "Record the migrated settlement count",
        "Execution_Scope": "TARGET_ONLY",
        "Comparison_Type": "NO_COMPARISON",
        "Result_Type": "INTEGER",
        "Target_Profile_Section": "target",
        "Target_Object": "dbo.Settlements",
        "Target_SQL": "SELECT COUNT(*) FROM dbo.Settlements",
        "Severity": "Low",
        "Owner": "migration-team",
        "Tags": "profile",
    },
    {
        "Test_ID": "TC-PAY-007",
        "Enabled": "No",
        "Domain": "Payments",
        "Flow": "Migration",
        "Test_Name": "Deferred until the refunds cutover",
        "Execution_Scope": "SOURCE_TARGET",
        "Comparison_Type": "EQUAL",
        "Result_Type": "INTEGER",
        "Source_Profile_Section": "source",
        "Source_SQL": "SELECT COUNT(*) FROM webservice.dbo.Refunds",
        "Target_Profile_Section": "target",
        "Target_SQL": "SELECT COUNT(*) FROM dbo.Refunds",
        "Severity": "Medium",
        "Owner": "migration-team",
        "Tags": "counts",
    },
)

#: Wording rules. Every one of these explains a status Python already decided.
DEMO_OBSERVATION_RULES: tuple[dict[str, str], ...] = (
    {
        "Enabled": "Yes",
        "Status": "PASS",
        "Platform": "ANY",
        "Error_Code": "ANY",
        # {actual_value} rather than both sides: a one-sided test has only one,
        # and "target ." reads like a bug rather than a scope.
        "Observation_Template": "{test_id} reconciled: {execution_scope} actual {actual_value}.",
    },
    {
        "Enabled": "Yes",
        "Status": "FAIL",
        "Platform": "COMPARISON",
        "Error_Code": "VALUE_MISMATCH",
        "Observation_Template": (
            "{test_id} mismatch: source {source_result} vs target {target_result} "
            "(variance {variance})."
        ),
    },
    {
        "Enabled": "Yes",
        "Status": "FAIL",
        "Platform": "COMPARISON",
        "Error_Code": "ANY",
        "Observation_Template": "{test_id} failed the {execution_scope} comparison. {error_detail}",
    },
    {
        "Enabled": "Yes",
        "Status": "ERROR",
        "Platform": "ANY",
        "Error_Code": "QUERY_TIMEOUT",
        "Observation_Template": (
            "{test_id} exceeded the {timeout_seconds}s timeout on {platform} ({database})."
        ),
    },
    {
        "Enabled": "Yes",
        "Status": "ERROR",
        "Platform": "ANY",
        "Error_Code": "ANY",
        "Observation_Template": "{test_id} errored on {platform}: {error_detail}",
    },
    {
        "Enabled": "Yes",
        "Status": "PROFILED",
        "Platform": "ANY",
        "Error_Code": "ANY",
        "Observation_Template": "{test_id} profiled {actual_value} (no pass/fail rule).",
    },
)

_EXECUTOR_CONTRACT: tuple[str, ...] = (
    "Executor contract (documentation only - nothing here is executed)",
    "",
    "1. TOML supplies every database connection. This workbook supplies none.",
    "2. Test Cases defines the queries and comparisons.",
    "3. Run Control defines execution behaviour.",
    "4. Observation Rules defines wording for a status Python has already decided.",
    "5. Comparison Types is validation metadata; comparisons are Python functions.",
    "6. The executor writes only the execution-output columns and Run History.",
    "7. Definition columns are never modified by a run.",
    "8. A disabled test never opens a connection and never runs SQL.",
    "9. Every query must be read-only and must return exactly one row and column.",
    "10. No password, token or secret is ever stored in TOML, in this workbook, or in a log.",
)

_CONVERSION_NOTES: tuple[str, ...] = (
    "Conversion notes (documentation only - nothing here is executed)",
    "",
    "Source SQL and target SQL are written in their own dialects and are passed",
    "to the driver verbatim. The executor never translates between dialects and",
    "never rewrites an identifier such as webservice.dbo, etables.dbo or",
    "DXBPRODSQL02.dbo. If an object cannot be reached, the login needs rights in",
    "that database, or the linked server needs defining - the SQL is not the",
    "thing to change.",
)

_CONNECTIONS_SHEET: tuple[str, ...] = (
    "Connections (example only - the executor never reads this sheet)",
    "",
    "Connection settings live in a TOML profile passed on the command line.",
    "This sheet exists so the shape of that file is visible next to the tests.",
    "",
    'version = "1.0"',
    "",
    "[workbook]",
    'path = "C:/Users/PC/Downloads/Payments.xlsx"',
    'sheet = "Test Cases"',
    "",
    "[source]",
    'type = "sqlserver"',
    'server = "localhost"',
    "port = 1433",
    'authentication = "windows"',
    "trust_server_certificate = true",
    'database = "webservice"',
    "",
    "[target]",
    'type = "sqlserver"',
    'server = "localhost"',
    "port = 1433",
    'database = "PaymentRecon_Target_Local"',
    "trust_server_certificate = true",
    'authentication = "windows"',
    "",
    "A password key is rejected, not ignored. Under password authentication the",
    "username lives in TOML and the password is typed at runtime, hidden.",
)


def build_workbook(
    tests: Sequence[Mapping[str, Any]] = DEMO_TESTS,
    *,
    run_control: Sequence[tuple[str, Any, str]] = DEFAULT_RUN_CONTROL,
    observation_rules: Sequence[Mapping[str, str]] = DEMO_OBSERVATION_RULES,
    include_documentation: bool = True,
) -> Workbook:
    """Build the whole template in memory."""
    workbook = Workbook()
    _build_test_cases(workbook, tests)
    _build_run_control(workbook, run_control)
    _build_observation_rules(workbook, observation_rules)
    _build_comparison_types(workbook)
    _build_run_history(workbook)
    if include_documentation:
        _build_text_sheet(workbook, "Executor Contract", _EXECUTOR_CONTRACT)
        _build_text_sheet(workbook, "Conversion Notes", _CONVERSION_NOTES)
        _build_text_sheet(workbook, "Connections", _CONNECTIONS_SHEET)
    return workbook


def write_workbook(path: str | Path, *, overwrite: bool = False, **kwargs: Any) -> Path:
    """Write the template to ``path``, refusing to clobber by accident."""
    destination = Path(path)
    if destination.exists() and not overwrite:
        raise FileExistsError(f"Refusing to overwrite existing file: {destination}")
    workbook = build_workbook(**kwargs)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    return destination


def _build_test_cases(workbook: Workbook, tests: Sequence[Mapping[str, Any]]) -> None:
    sheet = workbook.active
    sheet.title = TEST_CASES_SHEET
    headers = (*DEFINITION_COLUMNS, *OUTPUT_COLUMNS)
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.fill = _DEFINITION_FILL if header in DEFINITION_COLUMNS else _OUTPUT_FILL
        sheet.column_dimensions[get_column_letter(column)].width = _width(header)

    for offset, test in enumerate(tests):
        for column, header in enumerate(headers, start=1):
            if header in test:
                sheet.cell(row=2 + offset, column=column, value=test[header])
    sheet.freeze_panes = "A2"


def _build_run_control(workbook: Workbook, settings: Sequence[tuple[str, Any, str]]) -> None:
    sheet = workbook.create_sheet(RUN_CONTROL_SHEET)
    _write_table(sheet, ("Setting", "Value", "Description"), [list(row) for row in settings])
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 26
    sheet.column_dimensions["C"].width = 80


def _build_observation_rules(workbook: Workbook, rules: Sequence[Mapping[str, str]]) -> None:
    sheet = workbook.create_sheet(OBSERVATION_RULES_SHEET)
    headers = ("Enabled", "Status", "Platform", "Error_Code", "Observation_Template")
    _write_table(sheet, headers, [[rule.get(h, "") for h in headers] for rule in rules])
    for letter, width in zip("ABCDE", (10, 14, 14, 30, 90), strict=True):
        sheet.column_dimensions[letter].width = width


def _build_comparison_types(workbook: Workbook) -> None:
    """Documentation and validation metadata. The behaviour is in Python."""
    sheet = workbook.create_sheet(COMPARISON_TYPES_SHEET)
    rows = [
        [comparison.value, _COMPARISON_DESCRIPTIONS[comparison]] for comparison in ComparisonType
    ]
    _write_table(sheet, ("Comparison_Type", "Description"), rows)
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 90

    start = len(rows) + 3
    sheet.cell(row=start, column=1, value="Execution_Scope").font = _HEADER_FONT
    for offset, scope in enumerate(ExecutionScope, start=1):
        sheet.cell(row=start + offset, column=1, value=scope.value)
    sheet.cell(row=start, column=2, value="Result_Type").font = _HEADER_FONT
    for offset, result_type in enumerate(ResultType, start=1):
        sheet.cell(row=start + offset, column=2, value=result_type.value)

    status_column = 4
    sheet.cell(row=start, column=status_column, value="Status").font = _HEADER_FONT
    for offset, status in enumerate(TestStatus, start=1):
        sheet.cell(row=start + offset, column=status_column, value=status.value)
    sheet.cell(row=start, column=status_column + 1, value="Error_Code").font = _HEADER_FONT
    for offset, code in enumerate(ErrorCode, start=1):
        sheet.cell(row=start + offset, column=status_column + 1, value=code.value)
    sheet.column_dimensions[get_column_letter(status_column + 1)].width = 34


def _build_run_history(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(RUN_HISTORY_SHEET)
    _write_table(sheet, RUN_HISTORY_COLUMNS, [])
    for column in range(1, len(RUN_HISTORY_COLUMNS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 22


def _build_text_sheet(workbook: Workbook, title: str, lines: Sequence[str]) -> None:
    sheet = workbook.create_sheet(title)
    for line in lines:
        sheet.append([line])
    sheet.cell(row=1, column=1).font = _HEADER_FONT
    sheet.column_dimensions["A"].width = 96


def _write_table(sheet: Any, headers: Sequence[str], rows: Sequence[Sequence[Any]]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=header)
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.fill = _DEFINITION_FILL
    for offset, row in enumerate(rows):
        for column, value in enumerate(row, start=1):
            sheet.cell(row=2 + offset, column=column, value=value)
    sheet.freeze_panes = "A2"


_COMPARISON_DESCRIPTIONS: Mapping[ComparisonType, str] = {
    ComparisonType.EQUAL: "Source and target are equal after normalization.",
    ComparisonType.EQUAL_ABS_TOLERANCE: (
        "abs(target - source) <= Absolute_Tolerance. Numeric result types only."
    ),
    ComparisonType.EQUAL_PCT_TOLERANCE: (
        "The percentage variance against the source is within Percentage_Tolerance. "
        "A zero source passes only when the target is zero too, or an Absolute_Tolerance "
        "says how much drift from zero is acceptable."
    ),
    ComparisonType.EXPECTED_EQUAL: "The executed value equals Expected_Value.",
    ComparisonType.EXPECTED_ZERO: "The executed value is zero.",
    ComparisonType.LESS_THAN_OR_EQUAL: "The executed value is at most Expected_Value.",
    ComparisonType.GREATER_THAN_OR_EQUAL: "The executed value is at least Expected_Value.",
    ComparisonType.NON_ZERO: "The executed value is not zero.",
    ComparisonType.BOOLEAN_TRUE: "The executed value is TRUE. Result_Type must be BOOLEAN.",
    ComparisonType.TEXT_CASE_INSENSITIVE_EQUAL: (
        "Text values match ignoring case. Under SOURCE_TARGET the two sides are compared; "
        "otherwise the executed value is compared with Expected_Value."
    ),
    ComparisonType.NO_COMPARISON: "Record the value as PROFILED. Never reported as a pass.",
}

_WIDE_COLUMNS = frozenset({"Source_SQL", "Target_SQL", "Observation", "Error_Detail"})
_MEDIUM_COLUMNS = frozenset(
    {"Test_Name", "Source_Object", "Target_Object", "Executed_At_UTC", "Run_ID", "Error_Code"}
)


def _width(header: str) -> int:
    if header in _WIDE_COLUMNS:
        return 56
    if header in _MEDIUM_COLUMNS:
        return 26
    return 18
