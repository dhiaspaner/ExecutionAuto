"""Generates example workbooks from a schema.

Used to produce ``templates/reconciliation_template.xlsx`` and to build
throwaway workbooks in tests. Because the sheet is generated *from* the schema,
a template always agrees with the schema that describes it — including when the
schema uses different header text.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import WorkbookSchema

__all__ = ["DEMO_ROWS", "build_template_workbook", "write_template"]

_INPUT_FILL = PatternFill("solid", fgColor="DDEBF7")
_RESULT_FILL = PatternFill("solid", fgColor="FCE4D6")

#: Offline demonstration rows. The SQL is illustrative only — milestone 1 never
#: sends it to a database.
DEMO_ROWS: tuple[dict[str, Any], ...] = (
    {
        "test_case_id": "TC-PAY-001",
        "domain": "Payments",
        "entity": "Payment",
        "reconciliation_type": "Row count",
        "enabled": True,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": "SELECT COUNT(*) FROM PAYMENTS WHERE STATUS = 'POSTED'",
        "target_sql": "SELECT COUNT(*) FROM dbo.Payments WHERE Status = 'POSTED'",
        "comparison_rule": "equal",
        "tolerance": 0,
        "timeout_seconds": 120,
        "expected_result": "Counts match",
        "severity": "High",
    },
    {
        "test_case_id": "TC-PAY-002",
        "domain": "Payments",
        "entity": "Payment",
        "reconciliation_type": "Sum",
        "enabled": True,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": "SELECT NVL(SUM(AMOUNT), 0) FROM PAYMENTS",
        "target_sql": "SELECT COALESCE(SUM(Amount), 0) FROM dbo.Payments",
        "comparison_rule": "numeric_tolerance",
        "tolerance": "0.01",
        "timeout_seconds": 180,
        "expected_result": "Totals match within rounding",
        "severity": "High",
    },
    {
        "test_case_id": "TC-PAY-003",
        "domain": "Payments",
        "entity": "Payment",
        "reconciliation_type": "Orphan check",
        "enabled": True,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": (
            "SELECT COUNT(*) FROM PAYMENTS P "
            "WHERE NOT EXISTS (SELECT 1 FROM CUSTOMERS C WHERE C.ID = P.CUSTOMER_ID)"
        ),
        "target_sql": (
            "SELECT COUNT(*) FROM dbo.Payments P "
            "WHERE NOT EXISTS (SELECT 1 FROM dbo.Customers C WHERE C.Id = P.CustomerId)"
        ),
        "comparison_rule": "expected_zero",
        "tolerance": 0,
        "timeout_seconds": 120,
        "expected_result": "No orphaned payments",
        "severity": "Critical",
    },
    {
        "test_case_id": "TC-PAY-004",
        "domain": "Payments",
        "entity": "Refund",
        "reconciliation_type": "Row count",
        "enabled": True,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": "SELECT COUNT(*) FROM REFUNDS",
        "target_sql": "SELECT COUNT(*) FROM dbo.Refunds",
        "comparison_rule": "equal",
        "tolerance": 0,
        "timeout_seconds": 120,
        "expected_result": "Counts match",
        "severity": "Medium",
    },
    {
        "test_case_id": "TC-PAY-005",
        "domain": "Payments",
        "entity": "Duplicate",
        "reconciliation_type": "Duplicate check",
        "enabled": False,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": (
            "SELECT COUNT(*) FROM (SELECT REFERENCE FROM PAYMENTS "
            "GROUP BY REFERENCE HAVING COUNT(*) > 1)"
        ),
        "target_sql": (
            "SELECT COUNT(*) FROM (SELECT Reference FROM dbo.Payments "
            "GROUP BY Reference HAVING COUNT(*) > 1) d"
        ),
        "comparison_rule": "expected_zero",
        "tolerance": 0,
        "timeout_seconds": 120,
        "expected_result": "Deferred to the next run",
        "severity": "Low",
    },
    {
        "test_case_id": "TC-PAY-006",
        "domain": "Payments",
        "entity": "Ledger",
        "reconciliation_type": "Sum",
        "enabled": True,
        "source_type": "sqlserver",
        "source_connection": "STAGING_SQLSERVER",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": "SELECT COALESCE(SUM(Amount), 0) FROM stg.Ledger",
        "target_sql": "SELECT COALESCE(SUM(Amount), 0) FROM dbo.Ledger",
        "comparison_rule": "numeric_tolerance",
        "tolerance": "0.5",
        "timeout_seconds": 120,
        "expected_result": "Totals match within 0.50",
        "severity": "High",
    },
    {
        "test_case_id": "TC-PAY-007",
        "domain": "Payments",
        "entity": "Customer",
        "reconciliation_type": "Row count",
        "enabled": True,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": "SELECT COUNT(*) FROM CUSTOMERS",
        "target_sql": "SELECT COUNT(*) FROM dbo.Customers",
        "comparison_rule": "equal",
        "tolerance": 0,
        "timeout_seconds": 60,
        "expected_result": "Counts match",
        "severity": "High",
    },
    {
        "test_case_id": "TC-PAY-008",
        "domain": "Payments",
        "entity": "Settlement",
        "reconciliation_type": "Row count",
        "enabled": True,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": "SELECT COUNT(*) FROM SETTLEMENTS WHERE SETTLED_ON IS NOT NULL",
        "target_sql": "SELECT COUNT(*) FROM dbo.Settlements WHERE SettledOn IS NOT NULL",
        "comparison_rule": "equal",
        "tolerance": 0,
        "timeout_seconds": 120,
        "expected_result": "Counts match",
        "severity": "Critical",
    },
    {
        "test_case_id": "TC-PAY-009",
        "domain": "Payments",
        "entity": "Fee",
        "reconciliation_type": "Sum",
        "enabled": True,
        "source_type": "oracle",
        "source_connection": "LEGACY_ORACLE",
        "target_connection": "MIGRATED_SQLSERVER",
        "source_sql": "SELECT NVL(SUM(FEE), 0) FROM FEES",
        "target_sql": "SELECT COALESCE(SUM(Fee), 0) FROM dbo.Fees",
        "comparison_rule": "equal",
        "tolerance": 0,
        "timeout_seconds": 120,
        "expected_result": "Totals match",
        "severity": "Medium",
    },
)

_README_LINES: tuple[tuple[str, ...], ...] = (
    ("Migration reconciliation template",),
    (),
    ("This sheet is unrelated to execution and must survive a run untouched.",),
    ("It exists so the preserve_other_sheets guarantee is visible and testable.",),
    (),
    ("Column headers on the test-case sheet are bound to semantic fields by",),
    ("the schema file passed with --schema. Rename a header there, not in code.",),
)


def build_template_workbook(
    schema: WorkbookSchema,
    rows: Sequence[Mapping[str, Any]] = DEMO_ROWS,
    *,
    include_readme_sheet: bool = True,
) -> Workbook:
    """Build an in-memory workbook whose headers come from ``schema``."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = schema.sheet_name

    definitions = list(schema.fields.values())
    for column, definition in enumerate(definitions, start=1):
        cell = sheet.cell(row=schema.header_row, column=column, value=definition.header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.fill = _INPUT_FILL if definition.read else _RESULT_FILL
        sheet.column_dimensions[get_column_letter(column)].width = _column_width(definition.name)

    for offset, row_values in enumerate(rows):
        row_number = schema.first_data_row + offset
        for column, definition in enumerate(definitions, start=1):
            if not definition.read:
                continue
            if definition.name in row_values:
                sheet.cell(row=row_number, column=column, value=row_values[definition.name])

    sheet.freeze_panes = sheet.cell(row=schema.first_data_row, column=1)

    if include_readme_sheet:
        readme = workbook.create_sheet("Read Me")
        for line in _README_LINES:
            readme.append(list(line))
        readme.column_dimensions["A"].width = 80

    return workbook


def write_template(
    path: str | Path,
    schema: WorkbookSchema,
    rows: Sequence[Mapping[str, Any]] = DEMO_ROWS,
    *,
    include_readme_sheet: bool = True,
    overwrite: bool = False,
) -> Path:
    """Write a template workbook to ``path``.

    Refuses to clobber an existing file unless ``overwrite`` is explicit.
    """
    destination = Path(path)
    if destination.exists() and not overwrite:
        raise FileExistsError(f"Refusing to overwrite existing file: {destination}")
    workbook = build_template_workbook(schema, rows, include_readme_sheet=include_readme_sheet)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
    return destination


def _column_width(field_name: str) -> int:
    if field_name.endswith("_sql"):
        return 58
    if field_name in {"remarks", "expected_result"}:
        return 42
    if field_name.endswith(("_connection", "_at")):
        return 22
    return 18
